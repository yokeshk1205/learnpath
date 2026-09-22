from __future__ import annotations

import argparse
import importlib.util
import json
import re
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "LearnPath_Phase_Project_Report.docx"
BASE_PATH = ROOT / "scripts" / "build_learnpath_report.py"

spec = importlib.util.spec_from_file_location("learnpath_base_report", BASE_PATH)
base = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(base)
BASE_ALT = base.set_last_picture_alt

BLACK = "000000"
NAVY = "172554"
BLUE = "1D4ED8"
LIGHT_BLUE = "EFF6FF"
LIGHT_GRAY = "F3F4F6"
WHITE = "FFFFFF"

CURRENT_CHAPTER = 0
CURRENT_SECTION = 0
CURRENT_SUBSECTION = 0
TABLE_COUNTS: dict[int, int] = {}
TABLE_CAPTIONS: list[str] = []

KEEP_TABLES = {
    ("Approach", "Strength", "LearnPath position"),
    ("Layer", "Technology", "Responsibilities", "Not allowed to do"),
    ("Entity family", "Scope and key rule", "Examples of stored state"),
    ("Diagnostic classification", "Decision rule in plain language"),
    ("NetworkX output", "Purpose in LearnPath", "Decision boundary"),
    ("Model or baseline", "Validation NDCG@5", "Test NDCG@5", "Test ROC-AUC", "Test Precision@5"),
    ("Trust risk", "Implemented control", "Evidence a reviewer can inspect"),
    ("Verification area", "Recorded result", "Interpretation"),
    ("Future activity", "Technical purpose", "Acceptance evidence"),
}


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def cell_margins(cell, value=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def page_number_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def set_page_numbering(section, fmt: str, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), fmt)
    pg_num.set(qn("w:start"), str(start))


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.05)
    section.right_margin = Inches(0.85)
    section.different_first_page_header_footer = True
    set_page_numbering(section, "lowerRoman", 1)
    page_number_field(section.footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size, align in (
        ("Title", 18, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 12, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.alignment = align
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(8)

    for style_name in ("List Bullet", "List Number"):
        styles[style_name].font.name = "Times New Roman"
        styles[style_name].font.size = Pt(12)
        styles[style_name].paragraph_format.line_spacing = 1.35
        styles[style_name].paragraph_format.space_after = Pt(3)

    for name in ("Front Heading", "Cover Title", "Cover Meta", "Caption Text", "Contents Line"):
        if name not in styles:
            styles.add_style(name, 1)
    styles["Front Heading"].font.name = "Times New Roman"
    styles["Front Heading"].font.size = Pt(15)
    styles["Front Heading"].font.bold = True
    styles["Front Heading"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Cover Title"].font.name = "Times New Roman"
    styles["Cover Title"].font.size = Pt(20)
    styles["Cover Title"].font.bold = True
    styles["Cover Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Cover Meta"].font.name = "Times New Roman"
    styles["Cover Meta"].font.size = Pt(12)
    styles["Cover Meta"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Caption Text"].font.name = "Times New Roman"
    styles["Caption Text"].font.size = Pt(10)
    styles["Caption Text"].font.italic = True
    styles["Caption Text"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Contents Line"].font.name = "Times New Roman"
    styles["Contents Line"].font.size = Pt(11)
    styles["Contents Line"].paragraph_format.space_after = Pt(4)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_centered(doc: Document, text: str, size=12, bold=False, italic=False, before=0, after=8):
    p = doc.add_paragraph(style="Cover Meta")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, italic=italic)
    return p


def front_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Front Heading")
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(text.upper())
    set_font(r, size=15, bold=True)


def academic_paragraph(doc: Document, text: str, bold_prefix: str | None = None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.35)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(text[: len(bold_prefix)])
        set_font(r, bold=True)
        r2 = p.add_run(text[len(bold_prefix) :])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def academic_bullets(doc: Document, items: Iterable[str], level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(item), size=11.5)


def academic_numbered(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(item), size=11.5)


def clean_number(title: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\s+", "", title).strip()


def academic_section_heading(doc: Document, title: str, subtitle: str | None = None, new_page=True):
    global CURRENT_SECTION, CURRENT_SUBSECTION
    CURRENT_SECTION += 1
    CURRENT_SUBSECTION = 0
    label = f"{CURRENT_CHAPTER}.{CURRENT_SECTION} {clean_number(title).upper()}"
    p = doc.add_heading(level=2)
    set_font(p.add_run(label), size=12, bold=True)
    if subtitle:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(8)
        sp.paragraph_format.keep_with_next = True
        sp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_font(sp.add_run(subtitle), size=11, italic=True)
    return p


def academic_subheading(doc: Document, title: str, level=2):
    global CURRENT_SUBSECTION
    CURRENT_SUBSECTION += 1
    label = f"{CURRENT_CHAPTER}.{CURRENT_SECTION}.{CURRENT_SUBSECTION} {clean_number(title)}"
    p = doc.add_heading(level=3)
    set_font(p.add_run(label), size=12, bold=True)
    return p


def academic_page_break(doc: Document):
    doc.add_page_break()


def table_caption(headers: Sequence[str]) -> str:
    global TABLE_COUNTS
    TABLE_COUNTS[CURRENT_CHAPTER] = TABLE_COUNTS.get(CURRENT_CHAPTER, 0) + 1
    title = " and ".join(str(x) for x in (headers[0], headers[-1]) if x)
    title = re.sub(r"\s+", " ", title).strip()
    return f"Table {CURRENT_CHAPTER}.{TABLE_COUNTS[CURRENT_CHAPTER]} {title}"


def academic_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths=None, font_size=8.2):
    if tuple(str(value) for value in headers) not in KEEP_TABLES:
        for values in rows:
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.space_after = Pt(5)
            lead = str(values[0]).strip()
            set_font(p.add_run(lead + ". "), size=11.2, bold=True)
            details = []
            for header, value in zip(headers[1:], values[1:]):
                value_text = str(value).strip()
                if value_text:
                    details.append(f"{header}: {value_text}")
            set_font(p.add_run("; ".join(details) + ("." if details else "")), size=11.2)
        return None

    caption = table_caption(headers)
    TABLE_CAPTIONS.append(caption)
    cp = doc.add_paragraph(style="Caption Text")
    set_font(cp.add_run(caption), size=10, italic=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    mark_header_row(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, NAVY)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(str(header)), size=max(7.7, min(9.0, font_size)), bold=True, color=WHITE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_margins(cell)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), size=max(7.5, min(9.0, font_size)))
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cells[idx])
            if row_index % 2:
                shade(cells[idx], LIGHT_GRAY)
    if widths:
        total = sum(widths)
        scale = min(1.0, 6.5 / total)
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width * scale)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def academic_formula(doc: Document, label: str, formula: str, explanation: str):
    cp = doc.add_paragraph(style="Caption Text")
    set_font(cp.add_run(label), size=10, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(formula), name="Cambria Math", size=10.5, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing = 1.1
    p2.paragraph_format.space_after = Pt(8)
    set_font(p2.add_run(explanation), size=9, italic=True)


def patch_base_helpers() -> None:
    def academic_alt(doc: Document, description: str) -> None:
        BASE_ALT(doc, description)
        if "Diagnostic" in description:
            p = doc.add_paragraph(style="Caption Text")
            set_font(p.add_run("Figure 4.1 Evidence-driven adaptive diagnostic v5"), size=10, italic=True)

    base.add_paragraph = academic_paragraph
    base.add_bullets = academic_bullets
    base.add_numbered = academic_numbered
    base.add_section_heading = academic_section_heading
    base.add_subheading = academic_subheading
    base.add_table = academic_table
    base.add_formula_band = academic_formula
    base.page_break = academic_page_break
    base.set_last_picture_alt = academic_alt


def cover(doc: Document) -> None:
    add_centered(doc, "LEARNPATH", size=14, bold=True, before=18, after=34)
    p = doc.add_paragraph(style="Cover Title")
    p.paragraph_format.space_after = Pt(18)
    set_font(p.add_run("EVIDENCE-DRIVEN ADAPTIVE LEARNING USING\nKNOWLEDGE GRAPH CONSTRAINTS AND MACHINE LEARNING"), size=20, bold=True, color=NAVY)
    add_centered(doc, "FINAL PROJECT REPORT", size=15, bold=True, before=10, after=26)
    add_centered(doc, "Submitted by", size=12, italic=True, after=10)
    add_centered(doc, "LEARNPATH PROJECT TEAM", size=14, bold=True, after=34)
    add_centered(doc, "in partial fulfillment of the requirements for the award of the degree of", size=11, after=8)
    add_centered(doc, "BACHELOR OF ENGINEERING", size=13, bold=True, after=7)
    add_centered(doc, "in", size=11, after=7)
    add_centered(doc, "COMPUTER SCIENCE AND ENGINEERING", size=13, bold=True, after=40)
    add_centered(doc, "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", size=12, bold=True, after=8)
    add_centered(doc, "SEPTEMBER 2026", size=12, bold=True, before=22, after=0)
    doc.add_page_break()


def certificate(doc: Document) -> None:
    front_heading(doc, "Bonafide Certificate")
    academic_paragraph(doc, "This is to certify that the project report entitled “LearnPath: Evidence-Driven Adaptive Learning Using Knowledge Graph Constraints and Machine Learning” is a bona fide record of the work carried out by the LearnPath Project Team under project supervision, in partial fulfillment of the requirements for the award of the Bachelor of Engineering degree in Computer Science and Engineering during the academic year 2026–2027.")
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    for idx, (role, width) in enumerate((("PROJECT SUPERVISOR", 3.2), ("HEAD OF THE DEPARTMENT", 3.2))):
        cell = t.cell(0, idx)
        cell.width = Inches(width)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run("____________________________\n" + role), size=11, bold=True)
    add_centered(doc, "Submitted for the project viva-voce examination held on ____________________", size=11, before=70, after=60)
    t2 = doc.add_table(rows=1, cols=2)
    for idx, role in enumerate(("INTERNAL EXAMINER", "EXTERNAL EXAMINER")):
        p = t2.cell(0, idx).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run("____________________________\n" + role), size=11, bold=True)
    doc.add_page_break()


def abstract_page(doc: Document) -> None:
    front_heading(doc, "Abstract")
    text = (
        "LearnPath is an adaptive learning platform designed to determine what a learner should learn next and to explain that decision. The implemented system maintains one global learner-skill state across multiple course enrollments, identifies evidence-bounded skill gaps through an adaptive diagnostic, evaluates required prerequisite relationships with a confidence-aware knowledge graph, projects retention decay, generates safe candidate skills, and ranks only prerequisite-eligible candidates using a checked Random Forest model. The personalized path distinguishes recognized, current, recommended-next, upcoming, and locked skills. New practice or assessment evidence updates mastery and confidence, marks affected paths stale, and produces a new auditable path version after regeneration."
    )
    academic_paragraph(doc, text)
    academic_paragraph(doc, "The architecture deliberately separates responsibilities. PostgreSQL is the source of truth for curriculum, prerequisite edges, learner state, immutable evidence, and path history. The Express and TypeScript service owns authorization, mastery policy, prerequisite safety, candidate generation, and path orchestration. A Python FastAPI service contributes stateless Random Forest inference and NetworkX structural analytics. NetworkX calculates topological layers, downstream reach, betweenness, shortest foundation routes, bottlenecks, and counterfactual unlocks, but it cannot override a failed required prerequisite edge.")
    academic_paragraph(doc, "Because real longitudinal deployment data were unavailable, model evaluation used a deterministic curriculum-grounded simulation containing 1,000 learners and 20,000 candidate interactions. Learners were separated across training, validation, and test sets. The selected Random Forest achieved a held-out ROC-AUC of 0.801 and Precision@5 of 0.525; its NDCG@5 of 0.578 was slightly below the highest-skill-gap baseline at 0.581. These results support implementation feasibility and predictive signal under simulation, not a claim of proven real-world educational effectiveness. The report therefore combines algorithmic detail, implementation evidence, limitations, and a prospective validation plan.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    set_font(p.add_run("Keywords: "), bold=True)
    set_font(p.add_run("adaptive learning, diagnostic assessment, knowledge graph, prerequisite reasoning, mastery, retention, Random Forest, NetworkX, personalized learning path."), italic=True)
    doc.add_page_break()


def acknowledgement(doc: Document) -> None:
    front_heading(doc, "Acknowledgement")
    academic_paragraph(doc, "The LearnPath Project Team gratefully acknowledges the guidance of the project supervisor, the support of the Department of Computer Science and Engineering, and the feedback of faculty reviewers who encouraged the work to focus on an explainable adaptive learning loop rather than a generic learning-management interface.")
    academic_paragraph(doc, "We also acknowledge the open-source communities behind PostgreSQL, React, TypeScript, Express, FastAPI, scikit-learn, NetworkX, and the testing tools used in this implementation. Their software made it possible to build and evaluate a complete local system while retaining inspectable algorithms, reproducible manifests, and clear service boundaries.")
    academic_paragraph(doc, "Finally, we thank the learners and reviewers whose expected questions shaped the interface: what the learner knows, what is missing, why a skill is locked or unlocked, what should be learned next, why it was recommended, and how the path changes after new evidence.")
    doc.add_page_break()


TOC_ENTRIES = [
    ("CHAPTER 1", "INTRODUCTION"),
    ("1.1", "Project objective and novelty"),
    ("1.2", "Learner flow and adaptive loop"),
    ("CHAPTER 2", "LITERATURE SURVEY"),
    ("2.1", "Learner knowledge modeling"),
    ("2.2", "Adaptive diagnostic assessment"),
    ("2.3", "Educational graphs and learning paths"),
    ("2.4", "Retention and recommendation ranking"),
    ("2.5", "Comparative synthesis"),
    ("CHAPTER 3", "SYSTEM DESIGN"),
    ("3.1", "System architecture"),
    ("3.2", "Authoritative data model"),
    ("CHAPTER 4", "PROJECT DESCRIPTION"),
    ("4.1", "Module-wise implementation"),
    ("4.2", "Algorithms"),
    ("4.3", "Knowledge Graph v2"),
    ("4.4", "Machine learning lifecycle"),
    ("CHAPTER 5", "IMPLEMENTATION AND RESULT DISCUSSION"),
    ("5.1", "Trust and assurance"),
    ("5.2", "Database and API design"),
    ("5.3", "Demonstration scenarios"),
    ("5.4", "Verification and limitations"),
    ("CHAPTER 6", "CONCLUSION AND FUTURE WORK"),
    ("REFERENCES", "REFERENCES"),
    ("APPENDICES", "APPENDICES"),
]


def contents_pages(doc: Document, page_map: dict[str, int]) -> None:
    entries = TOC_ENTRIES
    for page_index, chunk in enumerate((entries[:13], entries[13:])):
        front_heading(doc, "Table of Contents")
        for number, title in chunk:
            key = f"{number} {title}" if number != title else title
            p = doc.add_paragraph(style="Contents Line")
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.25))
            prefix = "" if number == title else number + "  "
            dots = "." * max(4, 65 - len(prefix) - len(title))
            page = page_map.get(key, page_map.get(number, "—"))
            set_font(p.add_run(f"{prefix}{title} {dots} {page}"), size=11, bold=number.startswith("CHAPTER"))
        doc.add_page_break()


def list_of_tables(doc: Document, captions: Sequence[str]) -> None:
    chunks = [list(captions[i : i + 25]) for i in range(0, len(captions), 25)] or [[]]
    for chunk in chunks:
        front_heading(doc, "List of Tables")
        for caption in chunk:
            p = doc.add_paragraph(style="Contents Line")
            set_font(p.add_run(caption), size=10.5)
        doc.add_page_break()


def list_of_figures(doc: Document) -> None:
    front_heading(doc, "List of Figures")
    figures = [
        "Figure 1.1 LearnPath adaptive decision loop",
        "Figure 3.1 System architecture and decision authority",
        "Figure 4.1 Evidence-driven adaptive diagnostic v5",
        "Figure 4.2 Knowledge Graph v2 hybrid intelligence",
    ]
    for text in figures:
        p = doc.add_paragraph(style="Contents Line")
        set_font(p.add_run(text), size=11)
    doc.add_page_break()


def abbreviations(doc: Document) -> None:
    front_heading(doc, "List of Abbreviations")
    items = [
        ("API", "Application Programming Interface"),
        ("CAT", "Computerized Adaptive Testing"),
        ("DAG", "Directed Acyclic Graph"),
        ("JWT", "JSON Web Token"),
        ("KG", "Knowledge Graph"),
        ("ML", "Machine Learning"),
        ("NDCG", "Normalized Discounted Cumulative Gain"),
        ("P@5", "Precision at Five"),
        ("REST", "Representational State Transfer"),
        ("RF", "Random Forest"),
        ("ROC-AUC", "Receiver Operating Characteristic—Area Under Curve"),
        ("SQL", "Structured Query Language"),
        ("UI/UX", "User Interface and User Experience"),
    ]
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    mark_header_row(t.rows[0])
    for idx, heading in enumerate(("Abbreviation", "Expansion")):
        shade(t.cell(0, idx), NAVY)
        t.cell(0, idx).text = ""
        set_font(t.cell(0, idx).paragraphs[0].add_run(heading), size=10.5, bold=True, color=WHITE)
    for abbr, expansion in items:
        cells = t.add_row().cells
        for idx, value in enumerate((abbr, expansion)):
            cells[idx].text = ""
            set_font(cells[idx].paragraphs[0].add_run(value), size=10.5, bold=idx == 0)
            cell_margins(cells[idx], 110)
    doc.add_page_break()


def start_body(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.05)
    section.right_margin = Inches(0.85)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""
    page_number_field(section.footer.paragraphs[0])
    set_page_numbering(section, "decimal", 1)


def chapter(doc: Document, number: int, title: str) -> None:
    global CURRENT_CHAPTER, CURRENT_SECTION, CURRENT_SUBSECTION
    CURRENT_CHAPTER = number
    CURRENT_SECTION = 0
    CURRENT_SUBSECTION = 0
    if number != 1:
        doc.add_page_break()
    p = doc.add_heading(level=1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"CHAPTER {number}\n{title.upper()}"), size=14, bold=True)


def literature_survey(doc: Document) -> None:
    global CURRENT_SECTION, CURRENT_SUBSECTION
    topics = [
        ("Learner knowledge modeling", [
            "Bayesian Knowledge Tracing models a knowledge component using interpretable probabilities for initial knowledge, learning, guessing, and slipping [1]. Performance Factors Analysis instead uses prior opportunities, successes, and failures in a logistic formulation [2]. Deep Knowledge Tracing showed that recurrent neural networks can learn more complex interaction sequences [3]. These methods establish the importance of maintaining learner state over time.",
            "LearnPath adopts the stateful principle but does not claim a calibrated latent-state model because its available training evidence is synthetic rather than longitudinal and population representative. It therefore stores an inspectable mastery value, a separate confidence value, an evidence-state label, and every supporting observation. This conservative design creates a traceable migration path toward Bayesian or neural knowledge tracing when real data becomes available.",
        ]),
        ("Adaptive diagnostic assessment", [
            "Computerized adaptive testing chooses questions according to information value instead of presenting the same fixed test to every learner. Chang and Ying demonstrated that global information can improve item selection when the current ability estimate is uncertain [4]. LearnPath uses the same high-level idea but applies it to course skills and downstream path decisions rather than to a single population-comparable ability score.",
            "Its question-value policy combines discrimination, uncertainty, prerequisite-gateway value, decision impact, course coverage, retention risk, cognitive diversity, and misconception evidence. The test stops only after a minimum evidence budget and normally uses 16–22 questions, with an absolute maximum of 28. A single answer yields only a PROBED estimate. Stronger states require repeated, varied observations and successful application or analysis evidence.",
        ]),
        ("Educational graphs and learning paths", [
            "Prerequisite graphs represent which foundations are required before later skills. Liang et al. studied the recovery of concept prerequisite relations from course dependencies [6], while recent learning-path work combines graph representations with collaborative, sequential, or hybrid recommendation methods [7]. These studies motivate graph-aware recommendation but do not by themselves guarantee safe progression.",
            "LearnPath makes required prerequisite enforcement deterministic. PostgreSQL stores the directed acyclic curriculum graph and TypeScript evaluates every required edge with confidence-adjusted mastery. NetworkX supplies structural features and explanations but remains read-only. This prevents a statistical score or graph-centrality measure from overriding a missing foundation.",
        ]),
        ("Retention and recommendation ranking", [
            "Spaced-repetition research treats recall as a function of time and prior practice. Half-life regression showed that trainable forgetting models can improve recall prediction [8]. LearnPath begins with an explicit exponential retention projection because current data cannot support learner-specific or skill-specific decay calibration. The projection identifies revision candidates without overwriting stored mastery evidence.",
            "Random Forests provide nonlinear interactions, robustness to mixed signals, and inspectable feature importances [9]. LearnPath evaluates ranking with Precision@k and NDCG@k, which reward useful candidates near the top [10]. The ranker operates only after prerequisite eligibility is settled, making its probability a benefit estimate among safe choices rather than a permission to unlock content.",
        ]),
    ]
    for title, paragraphs in topics:
        CURRENT_SECTION += 1
        CURRENT_SUBSECTION = 0
        h = doc.add_heading(level=2)
        set_font(h.add_run(f"2.{CURRENT_SECTION} {title.upper()}"), size=12, bold=True)
        for text in paragraphs:
            academic_paragraph(doc, text)
    CURRENT_SECTION += 1
    h = doc.add_heading(level=2)
    set_font(h.add_run(f"2.{CURRENT_SECTION} COMPARATIVE SYNTHESIS"), size=12, bold=True)
    academic_table(doc, ["Approach", "Strength", "LearnPath position"], [
        ["Knowledge tracing", "Models proficiency over interactions", "Preserves auditable evidence and explicit confidence until real longitudinal calibration is possible"],
        ["Adaptive testing", "Spends questions where information gain is high", "Uses decision-focused skill evidence and bounded stopping instead of certifying complete knowledge"],
        ["Knowledge-graph recommendation", "Respects structural relationships", "Separates required-edge safety from bounded structural leverage"],
        ["Spaced repetition", "Responds to forgetting", "Projects retention to create revision candidates without altering historical mastery"],
        ["Machine-learned ranking", "Captures nonlinear benefit patterns", "Ranks only safe candidates and stores model/version provenance"],
    ], widths=[1.5, 2.2, 2.8], font_size=8.5)
    academic_paragraph(doc, "The literature therefore supports each component of the LearnPath loop, while the project’s engineering contribution lies in integrating them under explicit authority boundaries: evidence creates learner state, deterministic graph rules create the safe choice set, statistical inference ranks that set, and versioned regeneration makes the resulting adaptation visible and auditable.")


REFERENCES = [
    "A. T. Corbett and J. R. Anderson, “Knowledge tracing: Modeling the acquisition of procedural knowledge,” User Modeling and User-Adapted Interaction, vol. 4, no. 4, pp. 253–278, 1995, doi: 10.1007/BF01099821.",
    "P. I. Pavlik Jr., H. Cen, and K. R. Koedinger, “Performance Factors Analysis: A new alternative to knowledge tracing,” in Proc. 14th Int. Conf. Artificial Intelligence in Education, Brighton, U.K., 2009, pp. 531–538.",
    "C. Piech et al., “Deep knowledge tracing,” in Advances in Neural Information Processing Systems 28, 2015, pp. 505–513.",
    "H. H. Chang and Z. Ying, “A global information approach to computerized adaptive testing,” Applied Psychological Measurement, vol. 20, no. 3, pp. 213–229, 1996, doi: 10.1177/014662169602000303.",
    "A. A. Hagberg, D. A. Schult, and P. J. Swart, “Exploring network structure, dynamics, and function using NetworkX,” in Proc. 7th Python in Science Conf., Pasadena, CA, USA, 2008, pp. 11–16.",
    "C. Liang, J. Ye, Z. Wu, B. Pursel, and C. L. Giles, “Recovering concept prerequisite relations from university course dependencies,” in Proc. 31st AAAI Conf. Artificial Intelligence, 2017, doi: 10.1609/aaai.v31i1.10550.",
    "H. Ngo, K. Vo, and T. Nguyen, “Personalized learning path recommendations: Fusing knowledge graph embedding, sequence mining, and collaborative filtering,” in 2024 IEEE Int. Conf. Big Data, Washington, DC, USA, 2024, pp. 8145–8153.",
    "B. Settles and B. Meeder, “A trainable spaced repetition model for language learning,” in Proc. 54th Annual Meeting of the ACL, Berlin, Germany, 2016, pp. 1848–1858.",
    "L. Breiman, “Random forests,” Machine Learning, vol. 45, pp. 5–32, 2001, doi: 10.1023/A:1010933404324.",
    "K. Järvelin and J. Kekäläinen, “Cumulated gain-based evaluation of information retrieval techniques,” ACM Trans. Information Systems, vol. 20, no. 4, pp. 422–446, 2002.",
    "S. M. Lundberg and S. I. Lee, “A unified approach to interpreting model predictions,” in Advances in Neural Information Processing Systems 30, 2017, pp. 4765–4774.",
    "Y. Choi et al., “EdNet: A large-scale hierarchical dataset in education,” in Artificial Intelligence in Education, Springer, 2020, pp. 69–73.",
    "M. Mitchell et al., “Model cards for model reporting,” in Proc. Conf. Fairness, Accountability, and Transparency, 2019, pp. 220–229.",
]


def references(doc: Document) -> None:
    global CURRENT_CHAPTER
    doc.add_page_break()
    CURRENT_CHAPTER = 6
    p = doc.add_heading(level=1)
    set_font(p.add_run("REFERENCES"), size=14, bold=True)
    for idx, item in enumerate(REFERENCES, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.32)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        set_font(p.add_run(f"[{idx}] {item}"), size=10.5)


def future_work(doc: Document) -> None:
    global CURRENT_SECTION, CURRENT_SUBSECTION
    CURRENT_SECTION += 1
    CURRENT_SUBSECTION = 0
    h = doc.add_heading(level=2)
    set_font(h.add_run(f"6.{CURRENT_SECTION} FUTURE WORK"), size=12, bold=True)
    academic_paragraph(doc, "The next phase is empirical validation rather than adding more decorative features. A prospective study should compare the checked Random Forest policy with a highest-skill-gap baseline and a dependency-valid nonpersonalized path. Primary outcomes should be immediate post-assessment gain and delayed retention; secondary outcomes should include time to mastery, path acceptance, diagnostic length, and learner-reported clarity.")
    academic_table(doc, ["Future activity", "Technical purpose", "Acceptance evidence"], [
        ["Real diagnostic calibration", "Estimate item difficulty, discrimination, guessing, and misconception patterns", "Stable parameters and held-out response prediction"],
        ["Domain-expert graph review", "Validate required and recommended edges", "Reviewed DAG with documented edge provenance"],
        ["Prospective learner evaluation", "Measure educational effectiveness against transparent baselines", "Predeclared learning and retention outcomes"],
        ["Calibration and subgroup analysis", "Assess probability reliability and equitable behavior", "Confidence intervals, calibration curves, subgroup metrics"],
        ["Human-controlled model registry", "Support safe retraining and rollback", "Approved artifact, checksums, model card, rollback test"],
        ["Production hardening", "Improve privacy, accessibility, observability, and load resilience", "Institutional security review and measured service objectives"],
    ], widths=[1.55, 2.55, 2.4], font_size=8.3)


def replace_everywhere(doc: Document) -> None:
    replacements = {
        "Diagnostic v3": "Diagnostic v5",
        "diagnostic v3": "diagnostic v5",
        "Adaptive diagnostic v3": "Adaptive diagnostic v5",
        "adaptive diagnostic v3": "adaptive diagnostic v5",
        "fixed question budget": "bounded adaptive question budget",
        "in 15 questions": "within 12–28 questions, normally 16–22",
        "Fifteen questions": "A bounded diagnostic",
        "101 passing tests": "107 passing tests",
        "35 passing Python tests": "37 passing Python tests",
        "101 API and 41 web tests": "107 API and 41 web tests",
        "17 integration and 35 ML counts": "17 integration and 37 ML counts",
    }
    for p in doc.paragraphs:
        if not p.text:
            continue
        text = p.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        if text != p.text:
            p.text = text
            for run in p.runs:
                set_font(run, size=12 if not p.style.name.startswith("Caption") else 10)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text
                    for old, new in replacements.items():
                        text = text.replace(old, new)
                    if text != p.text:
                        p.text = text
                        for run in p.runs:
                            set_font(run, size=8.5)

    figure_map = {
        "Figure 1  LearnPath": "Figure 1.1 LearnPath",
        "Figure 2  The API": "Figure 3.1 The API",
        "Figure 3  Graph": "Figure 4.2 Graph",
    }
    diagnostic_assigned = False
    for p in doc.paragraphs:
        text = p.text.strip()
        for old, new in figure_map.items():
            if text.startswith(old):
                p.text = new + text[len(old):]
                p.style = doc.styles["Caption Text"]
                for run in p.runs:
                    set_font(run, size=10, italic=True)
        if not diagnostic_assigned and "Diagnostic" in text and text.startswith("Figure"):
            p.text = "Figure 4.1 Evidence-driven adaptive diagnostic v5"
            p.style = doc.styles["Caption Text"]
            diagnostic_assigned = True

    # Keep portrait figures within the printable body of the Word document.
    target_widths = {2: Inches(5.05), 3: Inches(4.75)}
    for index, shape in enumerate(doc.inline_shapes):
        if index in target_widths:
            ratio = shape.height / shape.width
            shape.width = target_widths[index]
            shape.height = int(shape.width * ratio)


def build_document(output: Path, page_map: dict[str, int], known_captions: Sequence[str]) -> list[str]:
    global CURRENT_CHAPTER, CURRENT_SECTION, CURRENT_SUBSECTION, TABLE_COUNTS, TABLE_CAPTIONS
    CURRENT_CHAPTER = CURRENT_SECTION = CURRENT_SUBSECTION = 0
    TABLE_COUNTS = {}
    TABLE_CAPTIONS = []
    patch_base_helpers()
    doc = Document()
    configure(doc)
    cover(doc)
    certificate(doc)
    abstract_page(doc)
    acknowledgement(doc)
    contents_pages(doc, page_map)
    list_of_tables(doc, known_captions)
    list_of_figures(doc)
    abbreviations(doc)
    start_body(doc)

    diagrams = {
        "adaptive": ROOT / "docs" / "report-assets" / "adaptive-loop.png",
        "architecture": ROOT / "docs" / "ieee-paper-assets" / "system-architecture.png",
        "diagnostic": ROOT / "docs" / "ieee-paper-assets" / "diagnostic-v5.png",
        "graph": ROOT / "docs" / "ieee-paper-assets" / "knowledge-graph-v2.png",
    }

    chapter(doc, 1, "Introduction")
    base.add_project_objective(doc)
    base.add_learner_flow(doc)
    doc.add_picture(str(diagrams["adaptive"]), width=Inches(6.45))
    BASE_ALT(doc, "LearnPath adaptive decision loop from learner evidence through path regeneration")
    cap = doc.add_paragraph(style="Caption Text")
    set_font(cap.add_run("Figure 1.1 LearnPath adaptive decision loop"), size=10, italic=True)

    chapter(doc, 2, "Literature Survey")
    literature_survey(doc)

    chapter(doc, 3, "System Design")
    base.add_architecture(doc, diagrams)
    base.add_data_model(doc)

    chapter(doc, 4, "Project Description")
    base.add_modules(doc)
    base.add_algorithms(doc, diagrams)
    base.add_knowledge_graph(doc, diagrams)
    base.add_ml_lifecycle(doc)

    chapter(doc, 5, "Implementation and Result Discussion")
    base.add_trust(doc)
    base.add_database_api(doc)
    base.add_scenarios(doc)
    base.add_verification(doc)

    chapter(doc, 6, "Conclusion and Future Work")
    base.add_conclusion(doc)
    future_work(doc)
    references(doc)

    doc.add_page_break()
    p = doc.add_heading(level=1)
    set_font(p.add_run("APPENDICES"), size=14, bold=True)
    CURRENT_CHAPTER = 7
    CURRENT_SECTION = 0
    base.add_appendix(doc)

    replace_everywhere(doc)
    core = doc.core_properties
    core.title = "LearnPath Final Project Report"
    core.subject = "Adaptive learning, diagnostic assessment, knowledge graph, machine learning and dynamic paths"
    core.author = "LearnPath Project Team"
    core.keywords = "LearnPath, adaptive learning, knowledge graph, diagnostic, mastery, retention, Random Forest, NetworkX"
    core.comments = "Structured to follow the supplied phase-report format while documenting the implemented LearnPath system."
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return list(TABLE_CAPTIONS)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--page-map", type=Path)
    args = parser.parse_args()
    page_map = json.loads(args.page_map.read_text(encoding="utf-8")) if args.page_map and args.page_map.exists() else {}
    provisional = args.output.with_name(args.output.stem + "_provisional.docx")
    captions = build_document(provisional, page_map, [])
    build_document(args.output, page_map, captions)
    if provisional.exists():
        provisional.unlink()
    print(args.output.resolve())
    print(f"tables={len(captions)}")


if __name__ == "__main__":
    main()
