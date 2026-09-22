from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "LearnPath_Detailed_Technical_Project_Report.docx"
ASSET_DIR = ROOT / "docs" / "report-assets"

NAVY = "111827"
BLUE = "4F46E5"
TEAL = "0F766E"
GREEN = "15803D"
AMBER = "B45309"
RED = "B91C1C"
SLATE = "475569"
LIGHT_BLUE = "EEF2FF"
LIGHT_TEAL = "ECFDF5"
LIGHT_AMBER = "FFFBEB"
LIGHT_GRAY = "F8FAFC"
MID_GRAY = "E2E8F0"
WHITE = "FFFFFF"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, bold=False, color=NAVY, size=8.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, header in enumerate(headers):
        set_cell_shading(header_row.cells[idx], NAVY)
        set_cell_text(header_row.cells[idx], header, bold=True, color=WHITE, size=8.2)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for col_idx, value in enumerate(values):
            set_cell_text(cells[col_idx], str(value), color=NAVY, size=font_size)
            if row_idx % 2 == 1:
                set_cell_shading(cells[col_idx], LIGHT_GRAY)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_paragraph(doc: Document, text: str, bold_prefix: str | None = None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        first, rest = text[: len(bold_prefix)], text[len(bold_prefix) :]
        r = p.add_run(first)
        r.bold = True
        p.add_run(rest)
    else:
        p.add_run(text)
    return p


def add_bullets(doc: Document, items: Iterable[str], level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.04
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        p.add_run(item)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def add_section_heading(doc: Document, title: str, subtitle: str | None = None, new_page=True):
    if new_page:
        page_break(doc)
    p = doc.add_heading(level=1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(23)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    if subtitle:
        sp = doc.add_paragraph(subtitle)
        sp.paragraph_format.space_after = Pt(12)
        sp.paragraph_format.keep_with_next = True
        sp.runs[0].font.size = Pt(10.5)
        sp.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(10)
    line.paragraph_format.keep_with_next = True
    line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━").font.color.rgb = RGBColor.from_string(BLUE)


def add_subheading(doc: Document, title: str, level=2):
    p = doc.add_heading(title, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 2 else 5)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_formula_band(doc: Document, label: str, formula: str, explanation: str):
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    set_cell_shading(table.cell(0, 0), LIGHT_BLUE)
    set_cell_text(table.cell(0, 0), label, bold=True, color=BLUE, size=8.5)
    cell = table.cell(1, 0)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(formula)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10)
    run.bold = True
    p2 = cell.add_paragraph(explanation)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    p2.runs[0].font.name = "Aptos"
    p2.runs[0].font.size = Pt(8)
    p2.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
    set_cell_margins(cell, top=120, bottom=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def set_last_picture_alt(doc: Document, description: str) -> None:
    inline = doc.inline_shapes[-1]._inline
    inline.docPr.set("descr", description)
    inline.docPr.set("title", description)


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/aptos.ttf"),
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    bold_candidates = [
        Path("C:/Windows/Fonts/aptos-bold.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf"),
    ]
    for candidate in bold_candidates if bold else candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def draw_arrow(draw, start, end, color=BLUE, width=6):
    draw.line([start, end], fill=f"#{color}", width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - direction * 16, y2 - 10), (x2 - direction * 16, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 10, y2 - direction * 16), (x2 + 10, y2 - direction * 16)]
    draw.polygon(pts, fill=f"#{color}")


def draw_box(draw, xy, title, lines, fill=WHITE, border=BLUE, title_color=NAVY):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=f"#{fill}", outline=f"#{border}", width=3)
    draw.text((x1 + 20, y1 + 15), title, font=font(25, True), fill=f"#{title_color}")
    y = y1 + 54
    for line in lines:
        draw.text((x1 + 20, y), line, font=font(17), fill=f"#{SLATE}")
        y += 25


def create_diagrams() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    outputs = {}

    # Adaptive loop
    img = Image.new("RGB", (1800, 650), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((50, 28), "The LearnPath adaptive decision loop", font=font(38, True), fill=f"#{NAVY}")
    labels = [
        ("Learner knowledge", "Global mastery\nconfidence evidence"),
        ("Skill gaps", "Course target minus\ncurrent knowledge"),
        ("Prerequisite gate", "Required edges first\nunknown never passes"),
        ("Candidate set", "Learn revision and\nsupporting foundations"),
        ("ML ranking", "Benefit probability\nfor eligible skills"),
        ("Personalized path", "Recognized current\nnext upcoming locked"),
    ]
    box_w, gap, y1, y2 = 255, 35, 120, 285
    for i, (title, lines) in enumerate(labels):
        x1 = 35 + i * (box_w + gap)
        fill_color = LIGHT_BLUE if i in (2, 4, 5) else LIGHT_GRAY
        draw_box(draw, (x1, y1, x1 + box_w, y2), title, lines.split("\n"), fill=fill_color)
        if i < len(labels) - 1:
            draw_arrow(draw, (x1 + box_w + 5, 202), (x1 + box_w + gap - 7, 202), width=5)
    draw_box(draw, (505, 395, 810, 570), "Learning and assessment", ["Lesson practice and", "server-scored evidence"], fill=LIGHT_TEAL, border=TEAL)
    draw_box(draw, (990, 395, 1295, 570), "Update and regenerate", ["Mastery retention", "stale path new version"], fill=LIGHT_AMBER, border=AMBER)
    draw_arrow(draw, (1430, 285), (1230, 395), color=AMBER)
    draw_arrow(draw, (990, 482), (815, 482), color=TEAL)
    draw_arrow(draw, (505, 482), (165, 295), color=TEAL)
    out = ASSET_DIR / "adaptive-loop.png"
    img.save(out, quality=95)
    outputs["adaptive"] = out

    # Architecture
    img = Image.new("RGB", (1700, 900), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((50, 28), "System architecture and decision authority", font=font(38, True), fill=f"#{NAVY}")
    draw_box(draw, (60, 110, 430, 330), "React and Vite client", ["Learner and reviewer views", "Renders server decisions", "No mastery or ranking logic"], fill=LIGHT_BLUE)
    draw_box(draw, (660, 95, 1040, 350), "Express TypeScript API", ["Authentication and validation", "Mastery and retention policy", "Prerequisite safety gates", "Candidate and path orchestration"], fill=LIGHT_TEAL, border=TEAL)
    draw_box(draw, (1270, 110, 1640, 330), "Python FastAPI service", ["Random Forest inference", "NetworkX graph analytics", "Checksum and schema checks", "Stateless read-only boundary"], fill=LIGHT_AMBER, border=AMBER)
    draw_box(draw, (660, 560, 1040, 815), "PostgreSQL source of truth", ["Curriculum and prerequisite graph", "Global learner skill state", "Append-only evidence ledger", "Versioned paths and feedback"], fill=LIGHT_GRAY, border=NAVY)
    draw_arrow(draw, (430, 220), (655, 220))
    draw_arrow(draw, (1045, 220), (1265, 220), color=AMBER)
    draw_arrow(draw, (850, 355), (850, 555), color=TEAL)
    draw_arrow(draw, (1265, 285), (1045, 595), color=AMBER)
    draw.text((95, 380), "Learner interaction", font=font(20, True), fill=f"#{BLUE}")
    draw.text((690, 405), "Authoritative business decisions", font=font(20, True), fill=f"#{TEAL}")
    draw.text((1300, 380), "Advisory intelligence", font=font(20, True), fill=f"#{AMBER}")
    out = ASSET_DIR / "architecture.png"
    img.save(out, quality=95)
    outputs["architecture"] = out

    # Diagnostic
    img = Image.new("RGB", (1700, 760), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((50, 28), "Evidence bounded adaptive diagnostic v3", font=font(38, True), fill=f"#{NAVY}")
    stages = [
        ("Coverage", "5 slots", ["Sample breadth", "Prefer high-value unknowns"]),
        ("Confirmation", "8 slots", ["Repeat uncertain skills", "Vary cognition and difficulty"]),
        ("Verification", "2 slots", ["Resolve contradictions", "Require application evidence"]),
        ("Result", "per skill", ["Interval and evidence state", "Untested remains explicit"]),
    ]
    positions = [(55, 140, 380, 390), (465, 140, 790, 390), (875, 140, 1200, 390), (1285, 140, 1610, 390)]
    fills = [LIGHT_BLUE, LIGHT_TEAL, LIGHT_AMBER, LIGHT_GRAY]
    borders = [BLUE, TEAL, AMBER, NAVY]
    for i, ((title, count, lines), pos) in enumerate(zip(stages, positions)):
        draw_box(draw, pos, title, [count] + lines, fill=fills[i], border=borders[i])
        if i < 3:
            draw_arrow(draw, (pos[2] + 8, 265), (positions[i + 1][0] - 8, 265), color=borders[i])
    draw.rounded_rectangle((310, 500, 1390, 680), radius=18, fill=f"#{LIGHT_GRAY}", outline=f"#{MID_GRAY}", width=3)
    draw.text((350, 525), "Key safety rule", font=font(26, True), fill=f"#{RED}")
    draw.text((350, 570), "One correct answer creates a PROBED estimate, never certified mastery.", font=font(25, True), fill=f"#{NAVY}")
    draw.text((350, 615), "MASTERED needs repeated varied observations plus successful APPLY or ANALYZE evidence.", font=font(21), fill=f"#{SLATE}")
    out = ASSET_DIR / "diagnostic.png"
    img.save(out, quality=95)
    outputs["diagnostic"] = out

    # Knowledge graph
    img = Image.new("RGB", (1700, 820), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((50, 28), "Knowledge Graph v2 hybrid intelligence", font=font(38, True), fill=f"#{NAVY}")
    draw_box(draw, (65, 115, 420, 330), "PostgreSQL graph", ["Skills and course mappings", "Required and recommended edges", "Single authoritative store"], fill=LIGHT_GRAY, border=NAVY)
    draw_box(draw, (505, 115, 860, 330), "TypeScript gate", ["Confidence-adjusted mastery", "Required edge enforcement", "Locked or eligible"], fill=LIGHT_TEAL, border=TEAL)
    draw_box(draw, (945, 115, 1300, 330), "NetworkX analytics", ["Layers reach and centrality", "Routes bottlenecks unlocks", "Read-only enrichment"], fill=LIGHT_AMBER, border=AMBER)
    draw_box(draw, (945, 505, 1300, 720), "Random Forest ranker", ["Eligible candidates only", "Benefit probability", "Graph bonus capped at 0.06"], fill=LIGHT_BLUE, border=BLUE)
    draw_box(draw, (505, 505, 860, 720), "Dependency-safe path", ["Recognized current next", "Upcoming and locked", "Reason codes retained"], fill=LIGHT_GRAY, border=NAVY)
    draw_arrow(draw, (425, 222), (500, 222), color=TEAL)
    draw_arrow(draw, (865, 222), (940, 222), color=AMBER)
    draw_arrow(draw, (1122, 335), (1122, 500), color=BLUE)
    draw_arrow(draw, (940, 612), (865, 612), color=BLUE)
    draw.text((70, 405), "Safety and eligibility are deterministic", font=font(22, True), fill=f"#{TEAL}")
    draw.text((1030, 405), "Structure informs priority but cannot unlock a skill", font=font(22, True), fill=f"#{AMBER}")
    out = ASSET_DIR / "knowledge-graph.png"
    img.save(out, quality=95)
    outputs["graph"] = out
    return outputs


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 34), ("Heading 1", 21), ("Heading 2", 14), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9)
    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(9)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "LEARNPATH  |  TECHNICAL PROJECT REPORT"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.name = "Aptos"
    hp.runs[0].font.size = Pt(7.5)
    hp.runs[0].font.bold = True
    hp.runs[0].font.color.rgb = RGBColor.from_string(SLATE)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("LearnPath  •  Panel evaluation edition  •  ")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(SLATE)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(64)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("LEARNPATH")
    run.font.name = "Aptos Display"
    run.font.size = Pt(15)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(14)
    tr = title.add_run("Detailed Technical\nProject Report")
    tr.font.name = "Aptos Display"
    tr.font.size = Pt(36)
    tr.bold = True
    tr.font.color.rgb = RGBColor.from_string(NAVY)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(28)
    sr = sub.add_run("AI based adaptive personalized learning path recommender")
    sr.font.name = "Aptos Display"
    sr.font.size = Pt(16)
    sr.font.color.rgb = RGBColor.from_string(SLATE)

    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(30)
    rr = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    rr.font.color.rgb = RGBColor.from_string(BLUE)

    add_table(
        doc,
        ["Report scope", "Coverage"],
        [
            ["System", "Architecture, modules, data model, APIs and learner workflow"],
            ["Intelligence", "Diagnostics, mastery, retention, prerequisites, NetworkX and ML ranking"],
            ["Assurance", "Evidence provenance, safety boundaries, evaluation, governance and limitations"],
        ],
        widths=[1.6, 5.3],
        font_size=8.6,
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(48)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(5)
    mr = meta.add_run("Prepared for project evaluation panel")
    mr.bold = True
    mr.font.size = Pt(11)
    mr.font.color.rgb = RGBColor.from_string(NAVY)
    meta2 = doc.add_paragraph("Version 1.0  |  September 2026")
    meta2.runs[0].font.size = Pt(10)
    meta2.runs[0].font.color.rgb = RGBColor.from_string(SLATE)


def add_executive_summary(doc: Document, diagrams: dict[str, Path]) -> None:
    add_section_heading(doc, "Executive summary", "What LearnPath does, why it is different, and what the implementation proves")
    add_paragraph(
        doc,
        "LearnPath is an adaptive learning system built to answer one operational question: what should this learner learn next, and why? It combines global learner knowledge, skill-gap analysis, prerequisite constraints, retention risk, machine-learned benefit ranking and versioned path regeneration. The application is not a generic learning-management interface with an isolated recommender. The recommendation loop is the organizing principle of the data model, API and learner experience.",
    )
    add_paragraph(
        doc,
        "The core architectural choice is the separation of authority from advice. PostgreSQL holds curriculum and learner truth. The TypeScript application API owns evidence updates, mastery, confidence, retention, prerequisite safety and path transactions. Python provides two bounded intelligence functions: NetworkX derives structural graph signals, and a checksum-validated Random Forest estimates recommendation benefit for candidates that have already passed the prerequisite gate. The browser only renders decisions returned by the backend.",
    )
    add_paragraph(
        doc,
        "The resulting path is evidence-aware and dependency-safe. It recognizes knowledge demonstrated in another course, exposes unknown or weak foundations, distinguishes revision from new learning, explains why a skill is locked or recommended, and becomes stale when later evidence changes the learner model. Regeneration creates a new immutable version rather than silently rewriting history.",
    )
    add_table(
        doc,
        ["Implemented scope", "Current project evidence"],
        [
            ["Curriculum", "5 compact courses and 36 shared active skills"],
            ["Learning coverage", "36 of 36 skills have lesson, diagnostic, practice and separate post-lesson assessment content"],
            ["Adaptive diagnostics", "Maximum 15 questions with coverage, confirmation and verification stages"],
            ["Recommendation model", "Random Forest, 55 frozen features, learner-disjoint evaluation, checked serving contract"],
            ["Demonstration", "5 synthetic learner stories covering cold start, uncertain knowledge, adaptation, retention and cross-course reuse"],
            ["Latest verification", "101 API tests and 41 web tests after Diagnostic v3; documented final audit also records 17 PostgreSQL integration and 35 ML tests"],
        ],
        widths=[1.55, 5.35],
    )
    doc.add_picture(str(diagrams["adaptive"]), width=Inches(6.95))
    set_last_picture_alt(doc, "LearnPath adaptive loop from learner knowledge through skill gaps, prerequisite gating, candidate generation, ML ranking, personalized path, learning, assessment and regeneration")
    cap = doc.add_paragraph("Figure 1  LearnPath closes the loop between evidence, safe recommendation and path adaptation")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)
    cap.runs[0].font.color.rgb = RGBColor.from_string(SLATE)


def add_contents(doc: Document) -> None:
    add_section_heading(doc, "Report contents", "A module-by-module technical and assurance narrative")
    rows = [
        ("1", "Project objective and novelty", "Problem, design principles and contribution"),
        ("2", "Learner flow and adaptive loop", "Visible journey from enrollment to regeneration"),
        ("3", "System architecture", "Client, API, database and intelligence service boundaries"),
        ("4", "Authoritative data model", "Global mastery, enrollment context, evidence and path history"),
        ("5", "Module-wise implementation", "Every major backend and frontend module"),
        ("6", "Algorithms", "Mastery, diagnostics, retention, graph, ranking and coordination"),
        ("7", "Knowledge Graph v2", "Hybrid graph architecture and novelty"),
        ("8", "Machine learning lifecycle", "Synthetic data, features, model selection and inference"),
        ("9", "Trust and assurance", "Why the system can be relied on and where it is limited"),
        ("10", "Database and API design", "Persistence guarantees and main endpoint families"),
        ("11", "Demonstration scenarios", "Five learner profiles and panel walkthrough"),
        ("12", "Verification and limitations", "Test evidence, completion audit and honest boundaries"),
        ("13", "Conclusion", "What the project demonstrates"),
        ("A", "Appendix", "Pseudocode, reason codes and panel defense points"),
    ]
    add_table(doc, ["Section", "Topic", "Purpose"], rows, widths=[0.65, 2.55, 3.7], font_size=8.6)
    add_subheading(doc, "Reading guide")
    add_bullets(
        doc,
        [
            "For the project idea and novelty, read Sections 1, 7 and 9.",
            "For an implementation defense, read Sections 3 through 8 and the Appendix.",
            "For the live panel demo, use Sections 2 and 11 as the presentation script.",
            "For limitations and responsible claims, use Sections 9 and 12.",
        ],
    )


def add_project_objective(doc: Document) -> None:
    add_section_heading(doc, "1 Project objective and novelty", "A system for deciding the next best learning action")
    add_subheading(doc, "Problem statement")
    add_paragraph(
        doc,
        "Traditional course platforms usually present the same module sequence to every learner. Completion percentages describe navigation, not knowledge. A learner may already know a prerequisite from another course, may have forgotten a previously mastered skill, or may understand an advanced topic while retaining a fragile basic foundation. A fixed sequence cannot represent these cases accurately.",
    )
    add_paragraph(
        doc,
        "LearnPath addresses the problem as a constrained decision system. The system must estimate what the learner knows, identify the gap relative to the active course, check prerequisite feasibility, rank only safe candidates, and update its decision when new evidence arrives. It must also show the learner and reviewer how that decision was made.",
    )
    add_subheading(doc, "Primary research and engineering contribution")
    add_table(
        doc,
        ["Contribution", "Implementation meaning", "Why it matters"],
        [
            ["Global Skill Passport", "One learner-skill state is reused across every course context", "Avoids relearning and demonstrates transfer of prior knowledge"],
            ["Evidence-bounded diagnostic", "Adaptive placement gathers repeated, varied evidence within a 15-question cap", "Avoids claiming mastery from one answer while keeping onboarding practical"],
            ["Hybrid Knowledge Graph v2", "Relational graph truth plus TypeScript safety gates and NetworkX structure analytics", "Combines auditability, safe prerequisites and graph intelligence without duplicate truth stores"],
            ["Constrained ML ranking", "Random Forest sees only prerequisite-eligible candidates", "ML optimizes priority but cannot make unsafe progression decisions"],
            ["Versioned adaptation", "Evidence marks an active path stale and explicit regeneration creates a new immutable version", "Makes change visible, explainable and auditable"],
            ["Cross-course coordination", "Shared skills are recognized while each enrollment keeps its own dependency-valid path", "Shows personalization across multiple simultaneous courses"],
        ],
        widths=[1.45, 3.0, 2.45],
    )
    add_subheading(doc, "Design principles")
    add_bullets(
        doc,
        [
            "Unknown is represented honestly. Missing evidence is not converted to zero knowledge or assumed mastery.",
            "Course completion and knowledge are separate. Passive activity cannot manufacture mastery.",
            "Safety precedes optimization. Required prerequisites are enforced before ML ranking.",
            "Every important decision is explainable from stored evidence, policy versions and reason codes.",
            "The system adapts only when learner evidence changes and preserves the earlier decision for comparison.",
            "Synthetic data and observational feedback are clearly labeled so the project does not overclaim real-world effectiveness.",
        ],
    )


def add_learner_flow(doc: Document) -> None:
    add_section_heading(doc, "2 Learner flow and adaptive loop", "How the implemented UI turns personalization into a visible workflow")
    add_numbered(
        doc,
        [
            "Authenticate and open the learner dashboard. The dashboard loads real enrollment, evidence, goal, retention and path summaries from the API.",
            "Enroll in one or more courses. Each enrollment establishes a course context but does not create duplicate skill mastery records.",
            "Take the course diagnostic when the starting evidence is insufficient. The server reveals one selected question at a time and stores answer drafts for resume.",
            "Submit the diagnostic. The server scores all responses, records immutable evidence, updates global skill estimates and returns tested, uncertain and explicitly untested skills.",
            "Generate the personalized path. The engine evaluates global mastery, retention and recursive prerequisite readiness, creates candidate lanes, requests benefit inference for eligible skills and persists a versioned path.",
            "Follow Learn Next directly into the lesson. Resource events record activity and time but do not change knowledge.",
            "Complete practice or a distinct assessment. Server-scored performance creates evidence, updates mastery and confidence, and may change retention or prerequisite eligibility.",
            "Observe the active path become stale. The UI explains that the learner model changed and offers explicit regeneration.",
            "Regenerate. The previous path becomes immutable SUPERSEDED history and a new ACTIVE path shows what moved, unlocked or became revision due.",
            "Switch courses. Shared global knowledge is recognized automatically while each course retains its own path and progress context.",
        ],
    )
    add_subheading(doc, "What the learner should be able to answer")
    add_table(
        doc,
        ["Learner question", "Where LearnPath answers it"],
        [
            ["What do I already know", "Global Skill Passport and recognized path lane"],
            ["What is uncertain", "Diagnostic results, mastery intervals, confidence and evidence state"],
            ["Why is this skill locked", "Prerequisite graph with exact required thresholds and shortfalls"],
            ["What should I learn next", "Current Learn Next path item and action"],
            ["Why was it selected", "Benefit score, gap, retention, course context, graph leverage and reason codes"],
            ["What changed after learning", "Before and after mastery, stale-path signal and immutable path comparison"],
        ],
        widths=[2.1, 4.8],
    )


def add_architecture(doc: Document, diagrams: dict[str, Path]) -> None:
    add_section_heading(doc, "3 System architecture", "Four layers with explicit decision ownership")
    doc.add_picture(str(diagrams["architecture"]), width=Inches(6.9))
    set_last_picture_alt(doc, "LearnPath system architecture showing the React client, authoritative Express API, PostgreSQL source of truth and advisory FastAPI intelligence service")
    cap = doc.add_paragraph("Figure 2  The API and database own learner truth; Python contributes bounded intelligence")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)
    cap.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
    add_table(
        doc,
        ["Layer", "Technology", "Responsibilities", "Not allowed to do"],
        [
            ["Presentation", "React, TypeScript, Vite", "Learner workflows, visual paths, graph, evidence and reviewer views", "Calculate mastery, unlock a skill or invent recommendation scores"],
            ["Application", "Express, TypeScript", "Authentication, validation, policy engines, orchestration and transactions", "Delegate prerequisite safety to the ML model"],
            ["Persistence", "PostgreSQL", "Curriculum, constraints, learner state, evidence, paths, feedback and migrations", "Store separate conflicting mastery per course"],
            ["Intelligence", "FastAPI, scikit-learn, NetworkX", "Checked benefit inference and read-only structural graph analytics", "Write learner data, satisfy required edges or rank locked skills"],
        ],
        widths=[1.0, 1.3, 2.8, 1.8],
    )
    add_subheading(doc, "Request and decision flow")
    add_paragraph(
        doc,
        "The browser sends an authenticated request to the Express API. The API loads a transactionally consistent learner and curriculum snapshot from PostgreSQL. Deterministic policies compute knowledge, retention and eligibility. Only the eligible candidate feature matrix is sent to FastAPI. The intelligence service validates the artifact checksum and the 55-field schema, returns probabilities and structural analytics, and never receives permission to mutate application state. The API applies bounded policy bonuses, orders the path dependency-safely and persists both the decision and its provenance.",
    )
    add_subheading(doc, "Failure behavior")
    add_bullets(
        doc,
        [
            "If PostgreSQL is unavailable, API readiness fails with HTTP 503 rather than serving incomplete state.",
            "If model metadata, schema or checksum is incompatible, inference returns an error; it does not silently substitute a heuristic score.",
            "If NetworkX analytics is unavailable, TypeScript prerequisite enforcement continues in safe mode and the graph bonus is omitted.",
            "A stale path is excluded from cross-course coordination until explicitly regenerated.",
        ],
    )


def add_data_model(doc: Document) -> None:
    add_section_heading(doc, "4 Authoritative data model", "Separating knowledge, context, activity and decisions")
    add_subheading(doc, "Core entities")
    add_table(
        doc,
        ["Entity family", "Scope and key rule", "Examples of stored state"],
        [
            ["Learner", "Authenticated person and profile", "Identity, session ownership and learner preferences"],
            ["Course enrollment", "One learner in one course context", "Status, module progress, diagnostic and path context"],
            ["Learner goal", "Optional relevance lens", "Goal skills and course links; dropping it preserves learning evidence"],
            ["Global learner skill", "Unique by learner_id and skill_id", "Mastery, confidence, evidence state and latest retention projection"],
            ["Skill evidence", "Append-only performance observations", "Source, score, difficulty, timing, before and after state, metadata"],
            ["Curriculum graph", "Shared active skills plus typed edges", "Course-skill mappings, REQUIRED and RECOMMENDED prerequisites"],
            ["Personalized path", "One current path per enrollment", "ACTIVE, STALE or SUPERSEDED version plus ordered items and provenance"],
            ["Recommendation feedback", "Response to the active Learn Next item", "Accepted or declined, reason, later assessed outcome and attribution"],
        ],
        widths=[1.35, 2.5, 3.05],
    )
    add_subheading(doc, "Knowledge state dimensions")
    add_table(
        doc,
        ["Dimension", "Meaning", "Important distinction"],
        [
            ["Mastery", "Estimated durable knowledge on a zero-to-one scale", "Not the same as lesson or module completion"],
            ["Confidence", "How strongly available evidence supports the estimate", "High mastery with weak confidence remains uncertain"],
            ["Evidence state", "UNKNOWN, ESTIMATED, ASSESSED or VERIFIED", "A policy label derived from volume, sessions, diversity and consistency"],
            ["Retention", "Current projected accessibility of demonstrated knowledge", "May fall with time while mastery remains the durable record"],
            ["Course progress", "Navigation and completion inside an enrollment", "Never copied into mastery"],
        ],
        widths=[1.2, 3.0, 2.7],
    )
    add_subheading(doc, "Evidence ledger guarantees")
    add_paragraph(
        doc,
        "All diagnostic, practice, quiz, assessment, module-assessment and retention-check observations pass through the same mastery service and create rows in skill_evidence. Database triggers reject update and delete operations. The ledger preserves the source event, measured score, correctness, difficulty, attempt number, timestamps, previous and resulting mastery, previous and resulting confidence, evidence-state transition, retention values and structured metadata. This is the audit trail behind every learner skill state.",
    )


MODULES = [
    (
        "Authentication and application security",
        "Provides registration, login, token refresh, logout and current-user resolution. Passwords are hashed, access tokens are signed, authentication endpoints are rate-limited, and protected routers use common middleware.",
        "Express routers validate inputs, issue short-lived access credentials, rotate refresh state and attach the authenticated learner identifier to the request. Authorization checks are repeated in enrollment, learning, practice and assessment services so a learner cannot operate on another learner's context.",
        "Authenticated learner identity and request metadata",
        "Authorized API context or explicit 4xx rejection",
    ),
    (
        "Catalog goals and enrollment",
        "Presents courses, course skills and optional learning goals; creates multiple active enrollments without duplicating learner knowledge.",
        "The catalog service reads active curriculum rows and goal mappings. The enrollment service creates the course relationship and module-progress records. Goal removal marks the goal dropped and unlinks context while preserving evidence, mastery and enrollment progress.",
        "Course, learner and optional goal selection",
        "Enrollment and context-specific progress state",
    ),
    (
        "Global Skill Passport",
        "Shows one reusable learner knowledge profile across every enrolled course.",
        "The learner-skills service joins global learner_skill_mastery with active skill and course mappings. It returns mastery, confidence, evidence state, retention, evidence counts and the courses that can reuse the skill. A database uniqueness constraint prevents course-specific duplicates.",
        "Learner identifier and optional filters",
        "Global skill records with reuse context",
    ),
    (
        "Central mastery and evidence engine",
        "Transforms verified performance events into bounded mastery, confidence and evidence-state transitions.",
        "Source-specific weights, difficulty-adjusted performance and evidence-history summaries are applied inside a row-level transaction. Confidence grows from volume, session count, source diversity, consistency and difficulty. Diagnostic confidence is capped at 0.45. The resulting state and immutable evidence are written atomically.",
        "Server-scored evidence and prior learner-skill state",
        "New mastery, confidence, evidence state and audit row",
    ),
    (
        "Adaptive diagnostic v3",
        "Estimates a defensible starting point without interrupting the learner with a long exhaustive exam.",
        "A server selector reveals one item at a time. Five coverage slots sample high-value breadth, eight confirmation slots deepen uncertain or contradictory observations, and two verification slots seek independent application evidence. Submission is all-or-nothing; classifications include MASTERED, READY, GAP, FORGOTTEN, FRAGILE_FOUNDATION, NEEDS_CONFIRMATION, PROBED and NOT_TESTED.",
        "Course or goal context, answer history and current learner state",
        "Evidence, intervals, classifications and path-generation handoff",
    ),
    (
        "Prerequisite engine",
        "Determines which skills are safe to learn now and explains every required edge.",
        "The service recursively loads course skills and supporting ancestors, orders the directed acyclic graph, and evaluates each REQUIRED edge against confidence-adjusted global mastery. RECOMMENDED edges remain advisory. Unknown mastery cannot satisfy a requirement. The output contains mastered, ready and locked nodes, shortfalls and reuse context.",
        "Enrollment or optional goal plus live global skill state",
        "Dependency order, readiness classifications and gate explanations",
    ),
    (
        "Retention and revision",
        "Detects when previously demonstrated knowledge is likely becoming inaccessible and schedules revision separately from new learning gaps.",
        "The retention service applies exponential decay from the latest immutable performance evidence. Confidence and evidence volume slow the effective decay rate. A skill becomes a revision candidate only if prior mastery was at least 0.65 and current retention is AT_RISK or CRITICAL. A retention check uses real server-scored questions and writes RETENTION_CHECK evidence.",
        "Mastery, confidence, evidence history and elapsed time",
        "Retention value, state and narrowly defined revision eligibility",
    ),
    (
        "Candidate generation",
        "Creates the complete decision space before machine learning.",
        "The engine combines course skills, recursive prerequisite ancestors, mastery gaps and revision needs. It assigns LEARN, REVISION or SUPPORTING_PREREQUISITE kinds and partitions results into eligible, locked and excluded lanes. Locked candidates retain explanations but never enter the inference matrix.",
        "Course context, global knowledge, retention and prerequisite result",
        "Eligible, locked and excluded candidate sets with reasons",
    ),
    (
        "Learning resources practice and assessment",
        "Turns recommendations into measurable learning actions.",
        "Lessons are mapped to skills and record activity events such as started, viewed and completed. Those events affect engagement only. Practice and distinct post-lesson assessment pools are selected by the server, correct options remain hidden until submission, and the mastery service receives scored performance evidence.",
        "Current path action, resource events or submitted answers",
        "Activity history or new performance evidence and state change",
    ),
    (
        "Synthetic research dataset",
        "Provides reproducible training and evaluation data without presenting synthetic learners as real users.",
        "A curriculum snapshot drives a seed-42 simulator with 1,000 synthetic learners and 20,000 interactions. It models mastery, prerequisites, retention, engagement, practice, completion, time, assessment improvement, feedback and explicit noise. Labels are determined from a weighted benefit score rather than assigned randomly. The data is never imported into the application database.",
        "Versioned active curriculum snapshot and fixed random seed",
        "Synthetic interactions, manifest, distributions and checksums",
    ),
    (
        "Feature engineering",
        "Converts decision-time learner and candidate context into a stable model contract.",
        "Feature version learner-candidate-features-v2 contains 55 numeric fields in a fixed order. Identifiers and labels remain separate. Current outcome fields are blocked, and historical outcome features are shifted by one interaction to prevent future leakage. A simulator-only latent ability variable was replaced with an observable performance proxy.",
        "Candidate interaction history available at recommendation time",
        "Normalized 55-feature matrix plus separate metadata and label",
    ),
    (
        "Model training and evaluation",
        "Selects a benefit-ranking model through reproducible learner-disjoint experimentation.",
        "Learners are split 70 percent training, 15 percent validation and 15 percent test with no overlap. Gradient Boosting, Random Forest and Logistic Regression are compared with Highest Skill Gap and Popularity baselines. Validation NDCG at 5 selects the model, ROC-AUC and F1 break ties, and the classification threshold is chosen on validation before the test set is opened.",
        "Frozen features, labels, split assignment and seed 42",
        "Selected Random Forest artifact, manifest, metrics and checksum",
    ),
    (
        "Checked inference service",
        "Serves model probabilities without allowing an incompatible or corrupt artifact to influence a path.",
        "FastAPI verifies artifact SHA-256, model version, feature version, inference version, feature count, exact field order and positive-class contract before prediction. Invalid bounds or schema receive 422; missing or incompatible artifacts receive 503. There is no heuristic fallback and the service is stateless.",
        "Eligible candidate records using the frozen feature contract",
        "Benefit probabilities with model and contract metadata",
    ),
    (
        "Personalized path engine",
        "Builds the visual centerpiece that connects prior knowledge, current action, future sequence and locked dependencies.",
        "The API combines inference probability with a 0.08 revision bonus and at most 0.06 times the NetworkX gateway score. It then applies dependency-aware ordering and stable tie-breaks. One path belongs to one enrollment and exposes RECOGNIZED, CURRENT, RECOMMENDED_NEXT, UPCOMING and LOCKED lanes. The stored snapshot includes decision-time mastery, confidence, retention, prerequisites, reason codes and version identifiers.",
        "Eligible candidates, ML scores, graph analytics and deterministic policy",
        "Versioned dependency-safe path and Learn Next explanation",
    ),
    (
        "Cross-course coordination",
        "Prevents multiple active courses from presenting confusing duplicate priorities while preserving each course path.",
        "Only current ACTIVE paths participate. The service groups identical global skills across course contexts and adjusts coordination priority using the best course priority, up to two additional active-course contexts and optional goal relevance. The result recommends a coordinated next action without merging or rewriting the underlying paths.",
        "Current active paths across the learner's enrollments",
        "Coordinated next action and cross-course reuse explanation",
    ),
    (
        "Dynamic adaptation and history",
        "Makes path changes explicit after meaningful learner evidence.",
        "New evidence marks affected ACTIVE paths STALE. Regeneration runs the complete prerequisite, graph, candidate, inference and ordering pipeline inside a transaction. The old row becomes SUPERSEDED, the new row becomes ACTIVE, and change summaries identify moved, added, removed, unlocked or revision-due skills.",
        "Stale path, current learner state and curriculum",
        "New immutable path version plus event and change summary",
    ),
    (
        "Recommendation feedback and attribution",
        "Captures learner response and later observed outcomes without claiming causality.",
        "Accepted or declined feedback applies only to the current Learn Next item and can include a structured rejection reason. Later assessed evidence within a 30-day window is compared with the recorded baseline. A change of at least plus or minus 0.05 is classified as improvement or decline; missing baseline remains OBSERVED_NO_BASELINE.",
        "Current recommendation response and later assessed evidence",
        "Version-linked observational outcome for evaluation",
    ),
    (
        "Analytics governance and reviewer tools",
        "Shows system performance honestly and blocks conclusions when the sample is too small.",
        "Learner analytics summarizes real backend state. Reviewer pages expose candidate partitions, synthetic data lineage, feature contracts, model metrics, inference checks and governance. Acceptance rates are withheld below 30 responses, calibration below 30 assessed outcomes, drift below 50 recent and 50 reference predictions, and retraining eligibility below 100 assessed recommendation outcomes. Promotion remains human-controlled.",
        "Stored path, feedback, prediction and outcome histories",
        "Metrics or explicit insufficient-sample status and governance decision",
    ),
    (
        "Modern learner interface",
        "Makes personalization immediately understandable rather than hiding it behind CRUD screens.",
        "React pages consume typed API clients and render enrollment-aware dashboard cards, diagnostic progress and results, the horizontal personalized path, Learn Next explanation, skill passport, prerequisite graph, retention states, lessons, practice, assessment, path history and cross-course recognition. Loading, empty, error and stale states are represented explicitly.",
        "Real API responses and route context",
        "Accessible learner actions and reviewer-visible intelligence",
    ),
]


def add_modules(doc: Document) -> None:
    add_section_heading(doc, "5 Module-wise implementation", "Inputs, internal logic and outputs of the implemented system")
    add_table(
        doc,
        ["Module family", "Primary responsibility", "Authority"],
        [
            ["Learner state", "Skill Passport, evidence, mastery, confidence and retention", "TypeScript plus PostgreSQL"],
            ["Placement", "Adaptive diagnostic and skill-gap classification", "TypeScript plus PostgreSQL"],
            ["Dependency intelligence", "Prerequisite enforcement and structural analytics", "TypeScript authoritative, NetworkX derived"],
            ["Recommendation", "Candidates, features, Random Forest rank and path ordering", "TypeScript orchestrated, Python advisory"],
            ["Learning loop", "Lessons, practice, assessment, feedback and regeneration", "TypeScript plus PostgreSQL"],
            ["Experience", "Learner and reviewer visualization", "React renders backend state"],
        ],
        widths=[1.35, 3.35, 2.2],
    )
    for index, (title, purpose, implementation, inputs, outputs) in enumerate(MODULES, start=1):
        if index in (5, 9, 13, 17):
            page_break(doc)
        add_subheading(doc, f"5.{index} {title}")
        add_paragraph(doc, purpose)
        add_paragraph(doc, "Technical implementation  " + implementation, bold_prefix="Technical implementation")
        add_table(doc, ["Inputs", "Outputs"], [[inputs, outputs]], widths=[3.45, 3.45], font_size=8.1)


def add_algorithms(doc: Document, diagrams: dict[str, Path]) -> None:
    add_section_heading(doc, "6 Algorithms", "How evidence becomes knowledge and knowledge becomes a path")
    add_subheading(doc, "6.1 Difficulty-aware mastery update")
    add_paragraph(
        doc,
        "Each server-scored observation becomes a bounded performance signal. The mastery service applies a source weight: 0.40 for assessment, diagnostic and module assessment, 0.30 for quiz, 0.25 for practice and 0.35 for retention check. Diagnostic reliability further reduces the effective weight when evidence is sparse or narrow.",
    )
    add_formula_band(
        doc,
        "Mastery update",
        "effective weight = source weight × reliability     |     mastery after = mastery before × (1 − effective weight) + measured performance × effective weight",
        "For a new skill, the diagnostic begins from a neutral 0.50 prior. All values are bounded to the valid zero-to-one interval.",
    )
    add_subheading(doc, "6.2 Confidence and evidence-state policy")
    add_paragraph(
        doc,
        "Confidence is not copied from mastery. It grows from independent evidence volume, multiple sessions, source diversity, answer consistency and difficulty coverage. Diagnostic evidence is capped at 0.45 because one placement session is not certification. ASSESSED requires at least five observations, two sessions and 0.45 confidence. VERIFIED requires at least twelve observations, four sessions, at least 0.70 consistency, at least 0.78 confidence, and source diversity or retention evidence.",
    )
    add_table(
        doc,
        ["State", "Interpretation", "Minimum policy meaning"],
        [
            ["UNKNOWN", "No performance observation", "Do not infer either weakness or mastery"],
            ["ESTIMATED", "Some evidence but still sparse or narrow", "Useful for personalization with visible uncertainty"],
            ["ASSESSED", "Repeated evidence meets assessment policy", "At least 5 observations, 2 sessions and confidence 0.45"],
            ["VERIFIED", "Strictest support for the estimate", "At least 12 observations, 4 sessions, consistency 0.70 and confidence 0.78 plus diversity"],
        ],
        widths=[1.0, 2.5, 3.4],
    )
    add_subheading(doc, "6.3 Adaptive diagnostic selection")
    doc.add_picture(str(diagrams["diagnostic"]), width=Inches(6.9))
    set_last_picture_alt(doc, "Diagnostic v3 sequence with five coverage slots, eight confirmation slots, two verification slots and evidence-bounded per-skill results")
    add_paragraph(
        doc,
        "The selector scores active course questions from signals available at that point in the session: unknown state, distance from target, confidence, retention risk, required downstream dependents, item discrimination, difficulty fit, cognitive novelty, missing application evidence and contradiction risk. Coverage favors different high-value skills. Confirmation selects a second or third independent observation. Verification resolves contradictions or missing application evidence. No skill receives more than three direct observations in one session.",
    )
    add_table(
        doc,
        ["Diagnostic classification", "Decision rule in plain language"],
        [
            ["MASTERED", "Three varied observations, correct application or analysis, target mastery and sufficient confidence"],
            ["READY", "Positive evidence supports proceeding but is below certification"],
            ["GAP", "Repeated observations indicate a missing foundation"],
            ["FORGOTTEN", "Prior strong mastery and retention risk conflict with weak current evidence"],
            ["FRAGILE_FOUNDATION", "Advanced performance is strong but a tested required prerequisite is weak"],
            ["NEEDS_CONFIRMATION", "Evidence conflicts or remains insufficient"],
            ["PROBED", "Exactly one direct observation"],
            ["NOT_TESTED", "No direct observation inside the fixed question budget"],
        ],
        widths=[1.75, 5.15],
    )
    add_subheading(doc, "6.4 Retention model")
    add_formula_band(
        doc,
        "Forgetting projection",
        "retention = mastery × exp(−effective decay × days since evidence)",
        "The base decay constant is 0.025. Stronger confidence and evidence volume reduce effective decay. STRONG begins at 0.75, AT RISK below 0.55 and CRITICAL below 0.30.",
    )
    add_paragraph(
        doc,
        "Retention is a current projection, not a retroactive rewrite of knowledge. Only a skill whose earlier mastery reached 0.65 can become a revision candidate when retention falls to AT_RISK or CRITICAL. This prevents an ordinary unlearned gap from being mislabeled as forgetting.",
    )
    add_subheading(doc, "6.5 Confidence-adjusted prerequisite gate")
    add_formula_band(
        doc,
        "Required-edge evaluation",
        "uncertainty penalty = (1 − confidence) × 0.10     |     effective mastery = max(0, observed mastery − uncertainty penalty)",
        "A REQUIRED edge passes only when effective mastery reaches the edge threshold. RECOMMENDED edges do not block. Missing mastery never passes.",
    )
    add_subheading(doc, "6.6 Candidate partition")
    add_numbered(
        doc,
        [
            "Load the enrollment's active course skills and recursively include required prerequisite ancestors.",
            "Join each skill to global learner mastery, confidence, evidence state and retention.",
            "Mark already strong reusable knowledge as recognized rather than relearning it.",
            "Create REVISION candidates only for demonstrated knowledge with at-risk or critical retention.",
            "Create SUPPORTING_PREREQUISITE candidates for foundations needed to open a course skill.",
            "Place candidates whose REQUIRED edges pass in eligible; keep failed candidates in locked with explanations; place out-of-scope or already-satisfied items in excluded.",
            "Build the ML feature matrix only from eligible candidates.",
        ],
    )
    add_subheading(doc, "6.7 Ranking and path priority")
    add_formula_band(
        doc,
        "Path priority",
        "priority = Random Forest benefit probability + 0.08 when revision is due + 0.06 × NetworkX gateway score",
        "Both bonuses are bounded policy adjustments. The graph score cannot exceed one, so its maximum effect is 0.06. Locked skills have no priority because they never reach inference.",
    )
    add_paragraph(
        doc,
        "The path builder then uses dependency level and stable course sequence or name tie-breaks. This produces deterministic ordering for equal priorities and ensures a supporting prerequisite appears before its dependent. The highest actionable item becomes RECOMMENDED_NEXT, while other eligible items become CURRENT or UPCOMING according to state and order.",
    )
    add_subheading(doc, "6.8 Multi-course coordination")
    add_formula_band(
        doc,
        "Coordination score",
        "min(1, best course priority + 0.04 × min(additional active course contexts, 2) + 0.06 × active-goal relevance)",
        "Only current ACTIVE paths participate. Grouping occurs by global skill, but each enrollment path remains independent and dependency-safe.",
    )
    add_subheading(doc, "6.9 Dynamic regeneration")
    add_paragraph(
        doc,
        "The state machine is ACTIVE to STALE to SUPERSEDED, with a newly generated ACTIVE version. Performance evidence that changes a relevant learner-skill state marks the path stale. Regeneration locks the current row, reruns the complete decision pipeline, supersedes the old path, inserts new items and writes a change summary in one transaction. At most one ACTIVE or STALE path exists per enrollment; history can contain multiple immutable SUPERSEDED versions.",
    )


def add_knowledge_graph(doc: Document, diagrams: dict[str, Path]) -> None:
    add_section_heading(doc, "7 Knowledge Graph v2", "The core novelty as a hybrid graph decision layer")
    doc.add_picture(str(diagrams["graph"]), width=Inches(5.8))
    set_last_picture_alt(doc, "Knowledge Graph v2 hybrid pipeline from PostgreSQL graph through TypeScript gates, NetworkX analytics, Random Forest ranking and a dependency-safe path")
    cap = doc.add_paragraph("Figure 3  Graph truth, safety, analytics and ranking remain separate concerns")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)
    cap.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
    add_subheading(doc, "Why PostgreSQL plus NetworkX")
    add_paragraph(
        doc,
        "PostgreSQL remains the only authoritative graph store because skills, course mappings and prerequisite edges are part of the same transactional curriculum as enrollments and learner state. Database constraints and cycle prevention already protect this graph. NetworkX receives a request-scoped, read-only projection for algorithms that are naturally expressed on a directed graph. This avoids a second synchronized truth store.",
    )
    add_paragraph(
        doc,
        "Neo4j was intentionally not introduced at this scale. It would be reasonable when the project requires very large heterogeneous graph traversal, independent graph teams, graph-native query workloads or cross-domain relationship discovery. For five courses and 36 shared skills, duplicating the curriculum in Neo4j would add synchronization, deployment and failure modes without improving the current learner decision. This is a deliberate architecture trade-off, not an absence of graph technology.",
    )
    add_subheading(doc, "Structural analytics")
    add_table(
        doc,
        ["NetworkX output", "Purpose in LearnPath", "Decision boundary"],
        [
            ["Topological layers", "Show dependency-valid learning levels", "Cannot change edge satisfaction"],
            ["Downstream reach", "Measure how many later skills depend on a foundation", "Used only as bounded leverage"],
            ["Betweenness centrality", "Identify bridge skills between graph regions", "Does not claim educational benefit by itself"],
            ["Shortest foundation route", "Explain a path from root knowledge to a target skill", "Read-only explanation"],
            ["Bottlenecks", "Find currently weak foundations that block many descendants", "TypeScript remains the source of lock status"],
            ["Counterfactual unlocks", "Show what would become ready if one prerequisite reached target", "Never writes hypothetical mastery"],
        ],
        widths=[1.5, 3.1, 2.3],
    )
    add_formula_band(
        doc,
        "Gateway score",
        "0.50 × descendant ratio + 0.30 × normalized betweenness + 0.20 × immediate counterfactual unlock ratio",
        "Every component is bounded from zero to one. The resulting graph contribution to final priority is capped at 0.06.",
    )
    add_subheading(doc, "How the learner sees the graph")
    add_bullets(
        doc,
        [
            "Mastered, ready and locked nodes use distinct visual states.",
            "REQUIRED and RECOMMENDED edges use different styles and explanations.",
            "Each locked node exposes observed mastery, confidence-adjusted mastery, target threshold and shortfall.",
            "The view explains shortest routes, downstream reach, bottlenecks and potential unlocks in plain language.",
            "Changing course context reprojects the same global learner knowledge onto a different portion of the graph.",
        ],
    )


def add_ml_lifecycle(doc: Document) -> None:
    add_section_heading(doc, "8 Machine learning lifecycle", "From explicit synthetic data to checked runtime inference")
    add_subheading(doc, "8.1 Data generation")
    add_table(
        doc,
        ["Property", "Implemented value"],
        [
            ["Classification", "SYNTHETIC / SIMULATED DATA"],
            ["Curriculum snapshot", "36 skills, 4-course source snapshot and 86 prerequisite edges in the training export"],
            ["Scale", "1,000 learners and 20,000 candidate interactions"],
            ["Random seed", "42"],
            ["Beneficial label", "Benefit score at or above 0.27"],
            ["Benefit weights", "0.40 mastery gain, 0.25 completion, 0.25 assessment improvement, 0.10 learner feedback"],
            ["Observed positive rate", "35.455 percent"],
            ["Separation", "Synthetic files remain under the ML data workspace and are not imported into PostgreSQL learner tables"],
        ],
        widths=[2.0, 4.9],
    )
    add_paragraph(
        doc,
        "The simulator models meaningful relationships rather than independent random columns. The stored manifest records positive prerequisite-mastery to post-assessment correlation, engagement to completion and mastery gain, practice to assessment improvement, and a negative relationship between difficulty mismatch and mastery gain. Noise flags include guessing, skipping, disengagement and inconsistent performance.",
    )
    add_subheading(doc, "8.2 Feature contract")
    add_paragraph(
        doc,
        "Version learner-candidate-features-v2 freezes 55 numeric features in exact order. The categories cover learner knowledge, evidence strength, recent performance, history, engagement, retention, prerequisite readiness, graph context, course context and candidate type. Learner, course, skill and interaction identifiers remain outside the numeric matrix. The target label is also excluded.",
    )
    add_table(
        doc,
        ["Leakage control", "Implementation"],
        [
            ["Decision-time only", "Features are calculated from information available before the candidate outcome"],
            ["Shifted history", "Completion, engagement, assessment improvement and feedback histories are shifted by one interaction"],
            ["Blocked fields", "Current completion, current engagement, current time, post-assessment, mastery gain, feedback, benefit score and label cannot enter features"],
            ["Observable proxy", "Simulator-only learner ability is removed and replaced with learner_performance_proxy"],
            ["Shared schema", "Training and inference load the same frozen feature names and order"],
        ],
        widths=[1.65, 5.25],
    )
    add_subheading(doc, "8.3 Experiment design")
    add_table(
        doc,
        ["Experiment element", "Policy"],
        [
            ["Split", "700 train learners and 14,000 rows; 150 validation learners and 3,000 rows; 150 test learners and 3,000 rows"],
            ["Isolation", "No learner appears in more than one split"],
            ["Candidates", "Gradient Boosting, Random Forest and Logistic Regression"],
            ["Baselines", "Highest Skill Gap and curriculum-derived Popularity"],
            ["Selection", "Highest validation NDCG at 5, with ROC-AUC and F1 tie-breaks"],
            ["Threshold", "Chosen by maximizing validation F1 over 181 grid points and then applied unchanged to test"],
            ["Test set", "Untouched during selection and opened only for final evaluation"],
        ],
        widths=[1.45, 5.45],
    )
    page_break(doc)
    add_subheading(doc, "8.4 Selected model")
    add_paragraph(
        doc,
        "The selected artifact is benefit-ranking-v2, a RandomForestClassifier with 220 estimators, maximum depth 12, square-root feature sampling, minimum leaf size 4, balanced class weighting, one execution job and seed 42. The classification threshold is 0.45. Important learned signals include recent score, current mastery, mastery gap, skill average score, historical feedback, score trend and historical completion rate.",
    )
    add_table(
        doc,
        ["Model or baseline", "Validation NDCG@5", "Test NDCG@5", "Test ROC-AUC", "Test Precision@5"],
        [
            ["Random Forest selected", "0.581021", "0.578463", "0.801410", "0.525333"],
            ["Gradient Boosting", "0.576017", "0.589501", "0.804062", "0.521333"],
            ["Logistic Regression", "0.564225", "0.580259", "0.802338", "0.513333"],
            ["Highest Skill Gap", "0.576408", "0.581229", "0.730785", "0.510667"],
            ["Popularity", "0.251677", "0.272412", "0.440132", "0.284000"],
        ],
        widths=[1.9, 1.25, 1.15, 1.15, 1.25],
        font_size=7.9,
    )
    add_paragraph(
        doc,
        "Interpretation  Random Forest was chosen strictly from validation NDCG@5, not from the more favorable test result of another model. On the held-out test, it improved ROC-AUC and Precision@5 over Highest Skill Gap but did not improve NDCG@5. This mixed result is important: it supports using the model as a bounded ranking component, not claiming that the synthetic experiment proves general superiority.",
        bold_prefix="Interpretation",
    )
    add_subheading(doc, "8.5 Deployment evolution and runtime contract")
    add_paragraph(
        doc,
        "The Phase 14 experiment manifest correctly records EVALUATED_NOT_DEPLOYED because ranking integration was outside that phase. Later implementation phases introduced the versioned FastAPI inference boundary and integrated the frozen selected artifact into path generation. Runtime loading verifies SHA-256 b67297304bd175b8a451c0a0f36011a8603b19d6fdf723a16e5ad94eb721fa6e, the model and feature versions, all 55 field names in order, the inference version and the positive class before deserialization and prediction.",
    )
    add_paragraph(
        doc,
        "The API never sends locked candidates to this service. Returned probability is therefore a priority estimate among safe options, not an eligibility decision. This distinction is central to the project's trust model.",
    )


def add_trust(doc: Document) -> None:
    add_section_heading(doc, "9 Trust and assurance", "How LearnPath supports confidence without overclaiming certainty")
    add_paragraph(
        doc,
        "Trust in LearnPath does not come from the phrase artificial intelligence. It comes from explicit boundaries, evidence provenance, conservative uncertainty handling, reproducible evaluation, failure-safe behavior and visible limitations. The system is designed so a reviewer can trace a recommendation from the path item back to eligibility checks, learner state, evidence and model version.",
    )
    add_table(
        doc,
        ["Trust risk", "Implemented control", "Evidence a reviewer can inspect"],
        [
            ["One answer treated as mastery", "Diagnostic one-answer result is PROBED; mastery needs repeated varied application evidence", "Attempt history, observation counts and result classification"],
            ["Unknown treated as knowledge", "Unknown mastery never satisfies REQUIRED prerequisite edges", "Gate explanation with missing evidence state"],
            ["ML unlocks an unsafe skill", "Locked skills removed before feature construction and inference", "Candidate eligible and locked lanes"],
            ["Course completion inflates knowledge", "Resource events update activity only and never call the mastery engine", "Evidence ledger contains only scored performance sources"],
            ["Evidence silently edited", "Append-only ledger with update and delete rejection triggers", "skill_evidence history and before or after state"],
            ["Corrupt or incompatible model", "Checksum, version, feature-order and class validation", "Inference readiness and artifact manifest"],
            ["Graph service unavailable", "Authoritative TypeScript gate continues and omits graph bonus", "Safe-mode analysis response"],
            ["Future leakage inflates metrics", "Decision-time allowlist, blocked fields, shifted histories and learner-disjoint splits", "Feature and experiment manifests"],
            ["Path history overwritten", "Stale then superseded immutable path version with transactional replacement", "Path history and change summary"],
            ["Feedback presented as causal", "Outcome labeled observational and requires a stored baseline", "Version-linked feedback and OBSERVED_NO_BASELINE state"],
            ["Small samples shown as reliable rates", "Minimum thresholds for rates, calibration, drift and retraining", "Governance page explains withheld metrics"],
        ],
        widths=[1.65, 3.25, 2.0],
        font_size=7.8,
    )
    add_subheading(doc, "Trust argument by layer")
    add_table(
        doc,
        ["Layer", "Trust claim", "Reasonable confidence boundary"],
        [
            ["Data", "Learner state is traceable", "Every update points to immutable scored evidence"],
            ["Policy", "Unsafe progression is blocked", "Required-edge decisions are deterministic and confidence-aware"],
            ["Graph", "Structural explanations are reproducible", "Analytics are derived from the current directed acyclic graph and never authoritative"],
            ["ML", "Inference is technically reproducible", "Artifact, features, split and metrics are versioned; educational effectiveness remains unproven on real learners"],
            ["Path", "A recommendation can be audited", "Snapshot retains scores, reasons, prerequisite state and model or policy versions"],
            ["UI", "Learner sees uncertainty and change", "Unknown, confidence, locked reasons, stale state and history are visible"],
        ],
        widths=[0.95, 2.5, 3.45],
    )
    page_break(doc)
    add_subheading(doc, "What the system can and cannot claim")
    add_table(
        doc,
        ["Supported claim", "Unsupported claim"],
        [
            ["The implemented pipeline produces dependency-safe personalized paths from real backend state", "The model is proven to improve learning outcomes in a production population"],
            ["Knowledge earned in one course is reused in another course that maps the same skill", "All domains share the same forgetting or mastery dynamics"],
            ["The diagnostic is more defensible than a single-question mastery judgment", "Fifteen questions can completely certify all 36 course skills"],
            ["Model evaluation is reproducible on the disclosed synthetic dataset", "Synthetic test performance is equivalent to external real-world validation"],
            ["Recommendation changes are explainable and versioned", "Observed improvement after acceptance is necessarily caused by the recommendation"],
        ],
        widths=[3.45, 3.45],
    )
    add_subheading(doc, "Validation roadmap")
    add_bullets(
        doc,
        [
            "Calibrate diagnostic items with real response data using item difficulty, discrimination and misconception patterns.",
            "Run a prospective learner study comparing learning gain, time to mastery and completion against a non-personalized sequence.",
            "Measure calibration and subgroup performance only after governance sample thresholds are met.",
            "Review prerequisite edges with domain experts and compare predicted unlocks with observed success.",
            "Retrain and promote only through a human-reviewed model registry process with rollback capability.",
        ],
    )


def add_database_api(doc: Document) -> None:
    add_section_heading(doc, "10 Database and API design", "Persistence guarantees and public service boundaries")
    add_subheading(doc, "Forward-only schema evolution")
    add_paragraph(
        doc,
        "The database is built through forward-only PostgreSQL migrations 0001 through 0028. The migration history evolves identity, curriculum, enrollments, diagnostics, mastery, evidence, prerequisites, retention, candidates, paths, feedback, governance, content coverage, Knowledge Graph v2 metadata and Diagnostic v3 question selection without destructive rebuilds.",
    )
    add_table(
        doc,
        ["Schema area", "Representative tables or constraints", "Guarantee"],
        [
            ["Curriculum", "courses, modules, skills, course_skills, skill_prerequisites", "Shared canonical skill identifiers and typed graph edges"],
            ["Learning context", "enrollments, enrollment_module_progress, learner_goals", "Progress belongs to enrollment; goals remain optional"],
            ["Knowledge", "learner_skill_mastery, skill_evidence", "Unique global state plus append-only provenance"],
            ["Assessment", "assessment_attempts, diagnostic_attempt_questions, diagnostic_answer_drafts, assessment_skill_results", "Ordered attempt snapshot, resume and final evidence summary"],
            ["Path", "personalized_paths, path_items, path_events and change summaries", "One current path per enrollment and immutable history"],
            ["Evaluation", "recommendation_feedback and attributed outcomes", "Feedback linked to exact path and model version"],
        ],
        widths=[1.25, 3.25, 2.4],
    )
    add_subheading(doc, "Primary endpoint families")
    add_table(
        doc,
        ["Family", "Representative endpoints", "Purpose"],
        [
            ["Authentication", "POST /register, /login, /refresh, /logout; GET /me", "Identity and session lifecycle"],
            ["Catalog and goals", "GET /overview, /goals/:id, /learner-goals; POST and DELETE learner goals", "Course discovery and optional relevance context"],
            ["Enrollment", "GET and POST /enrollments; status and module progress endpoints", "Multi-course ownership and progress"],
            ["Diagnostics", "GET /overview; POST /start; answer autosave, submit and results", "Adaptive placement flow"],
            ["Learner skills", "GET /learner-skills", "Global Skill Passport"],
            ["Prerequisites", "GET /prerequisites/analysis", "Dependency readiness and graph explanations"],
            ["Retention", "GET /retention", "Forgetting projection and revision state"],
            ["Candidates", "GET /:enrollmentId/candidates", "Eligible, locked and excluded lanes"],
            ["Paths", "GET path, history and coordination; POST generate and regenerate", "Versioned path lifecycle"],
            ["Learning and practice", "Resource overview, resource events, practice start and submit", "Learning actions and scored evidence"],
            ["Recommendations", "Evaluation and path feedback endpoints", "Response and later outcome attribution"],
            ["Analytics governance", "GET /analytics/overview and /governance/overview", "Learner state and sample-gated oversight"],
        ],
        widths=[1.3, 3.45, 2.15],
        font_size=7.8,
    )
    add_subheading(doc, "Transaction and concurrency model")
    add_bullets(
        doc,
        [
            "Mastery updates lock only the affected learner-skill row, limiting contention across independent skills.",
            "The scored evidence row and resulting learner-skill state are committed atomically.",
            "Path regeneration locks the current path and persists supersession, new path, items and change events as one operation.",
            "Uniqueness constraints prevent duplicate global mastery and more than one current ACTIVE or STALE path per enrollment.",
            "All endpoint inputs are schema-validated before service logic and authorization checks are scoped to the authenticated learner.",
        ],
    )


def add_scenarios(doc: Document) -> None:
    add_section_heading(doc, "11 Demonstration scenarios", "Five synthetic profiles that expose distinct system behavior")
    add_table(
        doc,
        ["Profile", "Starting state", "What to demonstrate", "Expected visible result"],
        [
            ["Maya Explorer", "No enrollments and unknown mastery", "Cold start and honest unknown state", "Five courses available; no invented mastery or path"],
            ["Noah Starter", "Enrolled with diagnostic incomplete", "Evidence gate before personalization", "Diagnostic is the next action; no premature personalized path"],
            ["Aisha Builder", "Functions learning flow", "Recommendation to evidence to regeneration", "Accepted lesson, practice, assessment, 40-point observed assessed gain, stale path and new version"],
            ["Elena Navigator", "Multiple courses and retention risk", "Cross-course coordination and revision", "Two stale paths, shared skills and retention review"],
            ["Ravi Strategist", "Advanced evidence and one focused gap", "Cross-course recognition and dependency focus", "20 recognized foundations with Basic Sorting isolated as the gap"],
        ],
        widths=[1.15, 1.7, 2.25, 1.8],
        font_size=7.9,
    )
    add_subheading(doc, "Recommended panel demonstration")
    add_numbered(
        doc,
        [
            "Open Noah to show why LearnPath does not create a path without enough starting evidence.",
            "Open Aisha and explain the current Learn Next card, including mastery gap, prerequisite state, model score and recommendation reason.",
            "Start the linked lesson. Point out that lesson completion changes engagement but not mastery.",
            "Complete practice and the separate assessment. Show the before and after mastery, confidence and evidence-state result returned by the server.",
            "Return to the path and show STALE rather than a silently changed sequence. Regenerate and compare the superseded and new versions.",
            "Open Elena or Ravi, switch course contexts and show that shared mastery is recognized globally while course progress and paths remain separate.",
            "Open the prerequisite graph to explain a locked node, its confidence-adjusted threshold and the counterfactual skills that would unlock.",
            "Finish in the reviewer model and governance views. Disclose synthetic training data, show the evaluation metrics and explain why rates remain withheld at small sample sizes.",
        ],
    )
    add_subheading(doc, "End-to-end Aisha trace")
    add_table(
        doc,
        ["Step", "Stored or computed change", "Visible proof"],
        [
            ["Recommendation", "Path item records skill, probability, policy bonuses, reasons and versions", "Learn Next explanation"],
            ["Acceptance", "Feedback linked to the active path item", "Accepted status and timestamp"],
            ["Lesson", "Resource activity and time recorded", "Learning status changes; mastery does not"],
            ["Practice", "PRACTICE evidence updates state with source weight 0.25", "Before and after mastery and confidence"],
            ["Assessment", "Separate ASSESSMENT evidence updates state with source weight 0.40", "Observed assessed gain and stronger evidence state"],
            ["Invalidation", "Affected current path becomes STALE", "Regenerate call to action"],
            ["Regeneration", "Old path SUPERSEDED and new path ACTIVE", "What moved, unlocked or changed"],
        ],
        widths=[1.0, 3.75, 2.15],
    )


def add_verification(doc: Document) -> None:
    add_section_heading(doc, "12 Verification and limitations", "Implemented evidence, honest boundaries and next validation work")
    add_subheading(doc, "Completion evidence")
    add_table(
        doc,
        ["Verification area", "Recorded result", "Interpretation"],
        [
            ["TypeScript checks", "Typecheck, ESLint and production API or web builds pass", "Static contracts and production compilation are healthy"],
            ["Latest API suite", "101 passing tests after Diagnostic v3", "Covers service and policy behavior including adaptive diagnostic safeguards"],
            ["Latest web suite", "41 passing tests after Diagnostic v3", "Covers typed API and learner component behavior"],
            ["Documented PostgreSQL suite", "17 passing integration tests in the final completion audit", "Verifies live database workflows and constraints"],
            ["Documented ML suite", "35 passing Python tests in the final completion audit", "Verifies data, feature, graph and inference contracts"],
            ["Content gate", "36 of 36 recommendable skills fully covered", "Every recommendation has lesson, diagnostic, practice and separate assessment content"],
            ["Browser workflows", "Five profiles and complete Aisha adaptation handoff passed", "Core UI flows were exercised against local services"],
        ],
        widths=[1.6, 2.25, 3.05],
        font_size=8.0,
    )
    add_paragraph(
        doc,
        "Test-count note  The final completion audit dated 1 September 2026 records 90 API and 40 web tests. Diagnostic v3 subsequently added coverage, bringing the latest TypeScript totals to 101 API and 41 web tests. The 17 integration and 35 ML counts are cited from that documented completion audit and are not presented as newly rerun in this report-generation step.",
        bold_prefix="Test-count note",
    )
    add_subheading(doc, "Known limitations")
    add_table(
        doc,
        ["Limitation", "Impact", "Responsible treatment"],
        [
            ["Synthetic model training", "Metrics may not transfer to real learners", "Always disclose synthetic lineage and require real prospective validation"],
            ["Compact curriculum", "Five courses and 36 skills do not represent every domain", "Treat graph and policy as extensible architecture, not universal content"],
            ["Fixed policy constants", "Mastery, decay and prerequisite thresholds require calibration", "Version constants and calibrate with domain and learner evidence"],
            ["Bounded diagnostic", "Not every skill can be directly tested in 15 questions", "Return NOT_TESTED and continue collecting evidence during learning"],
            ["Observational attribution", "Accepted recommendations and later gains do not prove causation", "Use associative labels and plan controlled evaluation"],
            ["Governance sample size", "Demo data cannot support stable rates, calibration or drift", "Withhold metrics below explicit minimum samples"],
            ["Local demonstration", "Operational scale and production resilience are not yet established", "Future work includes deployment hardening, monitoring and larger load tests"],
        ],
        widths=[1.5, 2.5, 2.9],
        font_size=7.9,
    )
    add_subheading(doc, "Acceptance criteria for a real deployment")
    add_bullets(
        doc,
        [
            "Expert review of curriculum skills, prerequisite edges and course mastery targets.",
            "Real response calibration for item difficulty, discrimination and guessing behavior.",
            "Prospective evaluation against a transparent baseline with predeclared learning-outcome measures.",
            "Calibration, subgroup and drift analysis after the minimum governance samples are reached.",
            "Security, privacy, data-retention and accessibility review appropriate to the institution.",
            "Human approval and rollback for any retrained model or policy version.",
        ],
    )


def add_conclusion(doc: Document) -> None:
    add_section_heading(doc, "13 Conclusion", "Strong core intelligence made visible to the learner")
    add_paragraph(
        doc,
        "LearnPath demonstrates a complete adaptive-learning architecture rather than a standalone prediction model. Its strongest contribution is the disciplined chain from evidence to knowledge, from knowledge to dependency-safe candidates, from candidates to bounded ML ranking, and from learner performance back to an explicitly regenerated path.",
    )
    add_paragraph(
        doc,
        "The Knowledge Graph v2 design is central to this result. PostgreSQL preserves one authoritative curriculum graph, TypeScript enforces safety, NetworkX adds structural explanation, and Random Forest estimates benefit only after eligibility is settled. Global mastery makes the graph useful across simultaneous courses, while retention and diagnostics keep the learner state current and evidence-bounded.",
    )
    add_paragraph(
        doc,
        "The project can be trusted as a transparent local demonstration because it preserves evidence, distinguishes uncertainty from mastery, validates model contracts, survives analytics failure safely, stores versioned decisions and refuses to turn small or synthetic samples into stronger claims than they support. The remaining research step is real-world educational validation, not hidden implementation work.",
    )
    add_table(
        doc,
        ["Panel takeaway", "Implemented proof"],
        [
            ["What the learner knows", "Global Skill Passport with mastery, confidence, evidence and retention"],
            ["What is missing", "Diagnostic classifications and skill-gap candidates"],
            ["Why something is locked", "Confidence-aware REQUIRED-edge explanation"],
            ["What to learn next", "Random-Forest-ranked eligible candidate in a dependency-safe path"],
            ["Why it was recommended", "Stored scores, graph leverage, retention state, course context and reason codes"],
            ["How the path changes", "Evidence marks the path stale and explicit regeneration creates an auditable new version"],
        ],
        widths=[2.0, 4.9],
    )


def add_appendix(doc: Document) -> None:
    add_section_heading(doc, "Appendix", "Technical pseudocode, reason vocabulary and panel defense points")
    add_subheading(doc, "A Adaptive path generation pseudocode")
    code_lines = [
        "INPUT enrollment_id, authenticated_learner_id",
        "LOAD enrollment, course skills, prerequisite ancestors and global learner-skill state",
        "PROJECT retention from latest immutable performance evidence",
        "EVALUATE REQUIRED prerequisite gates in TypeScript",
        "ANALYZE read-only graph projection with NetworkX when available",
        "BUILD candidates and partition into ELIGIBLE, LOCKED and EXCLUDED",
        "ASSERT no LOCKED candidate enters the feature matrix",
        "BUILD 55-feature records for ELIGIBLE candidates",
        "PREDICT benefit probability with checksum-validated Random Forest",
        "COMPUTE bounded revision and graph policy bonuses",
        "ORDER by dependency level, priority and stable tie-breaks",
        "PERSIST path version, items, decision-time snapshot and reason codes",
        "RETURN RECOGNIZED, CURRENT, RECOMMENDED_NEXT, UPCOMING and LOCKED lanes",
    ]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    set_cell_shading(table.cell(0, 0), LIGHT_GRAY)
    cell = table.cell(0, 0)
    cell.text = ""
    for i, line in enumerate(code_lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run(line)
        r.font.name = "Cascadia Mono"
        r.font.size = Pt(7.8)
        r.font.color.rgb = RGBColor.from_string(NAVY)
    add_subheading(doc, "B Diagnostic submission pseudocode")
    diagnostic_lines = [
        "LOCK the open attempt and verify ownership",
        "REJECT incomplete submissions and never expose correct options before submit",
        "SCORE each answer using the stored question snapshot",
        "GROUP observations by skill and calculate difficulty and discrimination weighted signals",
        "CLASSIFY each tested skill with repeated evidence and application constraints",
        "WRITE immutable evidence and update the global learner-skill row atomically",
        "RETURN mastery interval, before or after state, counts and explicit NOT_TESTED skills",
        "GENERATE or regenerate the enrollment path and open it directly",
    ]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    set_cell_shading(table.cell(0, 0), LIGHT_GRAY)
    cell = table.cell(0, 0)
    cell.text = ""
    for i, line in enumerate(diagnostic_lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        r = p.add_run(line)
        r.font.name = "Cascadia Mono"
        r.font.size = Pt(7.8)
        r.font.color.rgb = RGBColor.from_string(NAVY)
    add_subheading(doc, "C Core reason vocabulary")
    add_table(
        doc,
        ["Reason family", "Learner-facing meaning"],
        [
            ["Knowledge gap", "Current mastery remains below the course target"],
            ["Prerequisite satisfied", "Required foundations meet confidence-adjusted thresholds"],
            ["Supporting foundation", "This skill opens one or more dependent course skills"],
            ["Revision due", "Previously demonstrated knowledge has decayed to an actionable retention state"],
            ["Cross-course recognition", "Evidence earned elsewhere satisfies this course context"],
            ["High predicted benefit", "The selected model estimates this safe candidate is likely to help"],
            ["Locked prerequisite", "One or more REQUIRED edges remain below target or unknown"],
            ["Needs more evidence", "Current estimate is too uncertain for a stronger claim"],
        ],
        widths=[2.1, 4.8],
    )
    page_break(doc)
    add_subheading(doc, "D Short panel questions and answers")
    qa = [
        ("Where is machine learning used", "Only to rank prerequisite-eligible candidate skills by predicted benefit. It does not calculate mastery, unlock prerequisites or create learner evidence."),
        ("Why use NetworkX", "It provides topological layers, reach, centrality, routes, bottlenecks and counterfactual unlocks from a read-only graph projection."),
        ("Why not Neo4j", "At the current scale PostgreSQL already stores the constrained curriculum graph transactionally. A second graph database would duplicate truth and add synchronization risk without a demonstrated learner benefit."),
        ("Can one question prove mastery", "No. One answer is PROBED. Diagnostic mastery requires three varied observations, successful application or analysis evidence, target mastery and sufficient confidence."),
        ("How can fifteen questions test every skill", "They cannot and do not claim to. The adaptive budget gathers the highest-value evidence, returns untested skills explicitly and continues learning-state estimation through later practice and assessment."),
        ("What if a learner knows an advanced skill but forgot a basic one", "The diagnostic can classify FRAGILE_FOUNDATION, the prerequisite graph exposes the weak required edge, and the path recommends the supporting foundation without discarding advanced evidence."),
        ("How do you trust the model", "The artifact and feature contract are versioned and checksum-validated, evaluation uses learner-disjoint splits, leakage is blocked and the model is constrained behind prerequisite safety. Real-world effectiveness still requires prospective validation."),
        ("What is the core novelty", "An explainable adaptive loop built on global cross-course skill knowledge, evidence-bounded diagnostics, hybrid prerequisite graph intelligence, constrained ML ranking and auditable path regeneration."),
    ]
    add_table(doc, ["Question", "Answer"], qa, widths=[2.15, 4.75], font_size=8.1)
    add_subheading(doc, "E Source evidence inside the repository")
    add_bullets(
        doc,
        [
            "README.md for repository scope, local setup and learner routes.",
            "docs/architecture.md for authoritative boundaries and domain decisions.",
            "docs/diagnostic-v3.md for the evidence-bounded adaptive assessment policy.",
            "docs/knowledge-graph-v2.md for hybrid graph design and structural algorithms.",
            "docs/final-completion-audit.md and docs/demo-runbook.md for completion evidence and learner scenarios.",
            "services/ml/data and services/ml/models manifests for data lineage, feature contract, evaluation metrics and artifact checksum.",
            "database/migrations 0001 through 0028 and apps/api/src for executable persistence and application policies.",
        ],
    )


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    diagrams = create_diagrams()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_executive_summary(doc, diagrams)
    add_contents(doc)
    add_project_objective(doc)
    add_learner_flow(doc)
    add_architecture(doc, diagrams)
    add_data_model(doc)
    add_modules(doc)
    add_algorithms(doc, diagrams)
    add_knowledge_graph(doc, diagrams)
    add_ml_lifecycle(doc)
    add_trust(doc)
    add_database_api(doc)
    add_scenarios(doc)
    add_verification(doc)
    add_conclusion(doc)
    add_appendix(doc)
    core = doc.core_properties
    core.title = "LearnPath Detailed Technical Project Report"
    core.subject = "Architecture, modules, algorithms, machine learning, knowledge graph and trust assurance"
    core.author = "LearnPath Project Team"
    core.keywords = "LearnPath, adaptive learning, knowledge graph, diagnostic, Random Forest, NetworkX"
    core.comments = "Generated from the implemented LearnPath repository and its versioned manifests."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
