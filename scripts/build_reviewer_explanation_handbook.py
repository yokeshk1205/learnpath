from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "LearnPath_Reviewer_Explanation_Handbook.docx"
ASSETS = ROOT / "docs" / "reviewer-handbook-assets"
EQUATIONS = ASSETS / "equations"
REPORT_ASSETS = ROOT / "docs" / "report-assets"

BLACK = "000000"
NAVY = "172033"
BLUE = "3347B0"
TEAL = "0B766E"
AMBER = "A35409"
RED = "A61B1B"
WHITE = "FFFFFF"
TEXT = "1F2937"
MUTED = "536174"
PALE_BLUE = "F2F5FF"
PALE_TEAL = "F0FAF8"
PALE_AMBER = "FFF8EB"
PALE_GRAY = "F7F8FA"
BORDER = "D9D9D9"


def set_run_font(run, name="Aptos", size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), BORDER)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell(cell, value: str, *, bold=False, color=TEXT, size=9.0, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.03
    if align is not None:
        p.alignment = align
    set_run_font(p.add_run(str(value)), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths=None, size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_borders(table)
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], NAVY)
        set_cell(table.rows[0].cells[index], header, bold=True, color=WHITE, size=8.8)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            if row_index % 2:
                shade(cells[col_index], PALE_GRAY)
            set_cell(cells[col_index], value, size=size)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    return table


def para(doc: Document, text: str, *, lead: str | None = None, italic=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.14
    p.paragraph_format.keep_together = keep
    if lead and text.startswith(lead):
        set_run_font(p.add_run(lead), bold=True, size=10.5, color=BLACK)
        set_run_font(p.add_run(text[len(lead):]), size=10.5, color=TEXT, italic=italic)
    else:
        set_run_font(p.add_run(text), size=10.5, color=TEXT, italic=italic)
    return p


def bullets(doc: Document, values: Iterable[str], level=0):
    for value in values:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.09
        set_run_font(p.add_run(value), size=10.2, color=TEXT)


def numbered(doc: Document, values: Iterable[str]):
    for value in values:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.09
        set_run_font(p.add_run(value), size=10.2, color=TEXT)


def heading(doc: Document, text: str, level=1, new_page=False):
    if new_page:
        doc.add_page_break()
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        set_run_font(run, name="Aptos Display", size={1: 23, 2: 15, 3: 11.5}.get(level, 10.5), bold=True, color=BLACK)
    return p


def panel_words(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    set_run_font(p.add_run("How to explain it to the panel  "), bold=True, size=10.3, color=BLACK)
    set_run_font(p.add_run(text), size=10.3, color=MUTED, italic=True)


def code_location(doc: Document, files: Sequence[str]):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run("Where it is implemented  "), bold=True, size=9.5, color=BLACK)
    set_run_font(p.add_run("  |  ".join(files)), name="Cascadia Mono", size=8.3, color=MUTED)


def set_picture_alt(doc: Document, description: str):
    inline = doc.inline_shapes[-1]._inline
    inline.docPr.set("descr", description)
    inline.docPr.set("title", description)


def add_picture(doc: Document, path: Path, width: float, caption: str, alt: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt(doc, alt)
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(9)
    set_run_font(c.add_run(""), size=8)
    for run in c.runs:
        set_run_font(run, size=8.2, color=MUTED, italic=True)


def equation(doc: Document, name: str, width: float, caption: str, definitions: str):
    add_picture(doc, EQUATIONS / f"{name}.png", width, caption, f"Equation for {caption}")
    para(doc, definitions, lead="Terms  ")


def configure(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.14
    for name, size in (("Title", 32), ("Heading 1", 23), ("Heading 2", 15), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.size = Pt(10.2)
        doc.styles[name].font.color.rgb = RGBColor.from_string(TEXT)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("LearnPath reviewer explanation handbook   "), size=7.8, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)


def font(size, bold=False):
    choices = [
        Path("C:/Windows/Fonts/aptos-bold.ttf") if bold else Path("C:/Windows/Fonts/aptos.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf") if bold else Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf") if bold else Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in choices:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def draw_arrow(draw, start, end, color=BLUE, width=5):
    draw.line([start, end], fill=f"#{color}", width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        d = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - d * 15, y2 - 9), (x2 - d * 15, y2 + 9)]
    else:
        d = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 9, y2 - d * 15), (x2 + 9, y2 - d * 15)]
    draw.polygon(points, fill=f"#{color}")


def draw_box(draw, xy, title, lines, fill=PALE_GRAY, border=NAVY):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=f"#{fill}", outline=f"#{border}", width=3)
    draw.text((x1 + 18, y1 + 14), title, font=font(23, True), fill=f"#{BLACK}")
    y = y1 + 52
    for line in lines:
        draw.text((x1 + 18, y), line, font=font(17), fill=f"#{MUTED}")
        y += 25


def create_equation_images():
    """Render readable equation plates without relying on a system LaTeX install."""
    EQUATIONS.mkdir(parents=True, exist_ok=True)
    formulas = {
        "mastery": "w_eff = w_source x r    ;    M_new = clip[0,1](M_old(1 - w_eff) + P x w_eff)",
        "confidence": "C = clip[0,1](0.08 + 0.45(1 - exp(-n/6)) + 0.05 min(D,3) + 0.12 min(S/4,1) + 0.06 min(d_avg/5,1) + 0.12K + 0.10R)",
        "diagnostic_reliability": "r_diag = min(0.78, 0.16 min(n,3) + 0.06 min(L,3) + 0.08A + 0.06K)",
        "mastery_interval": "I = [clip(M - u), clip(M + u)]    ;    u = 0.06 + 0.22(1 - C) / sqrt(max(1,n))",
        "retention": "lambda_eff = 0.025(1.15 - 0.5C) / (1 + min(0.12 ln(1+n),0.45))    ;    R = M exp(-lambda_eff t)",
        "prerequisite": "p = 0.10(1 - C)    ;    M_eff = max(0, M - p)    ;    gate passes iff M_eff >= T",
        "gateway": "G = 0.50 D_r + 0.30 B_n + 0.20 U_r",
        "priority": "Q = min(1, p_benefit + 0.08 I_revision + 0.06G)",
        "coordination": "Q_c = min(1, Q_best + 0.04 min(k - 1, 2) + 0.06g)",
    }
    for name, text in formulas.items():
        canvas = Image.new("RGB", (1900, 190), "white")
        draw = ImageDraw.Draw(canvas)
        size = 38
        face = font(size, True)
        while draw.textbbox((0, 0), text, font=face)[2] > 1780 and size > 22:
            size -= 1
            face = font(size, True)
        bounds = draw.textbbox((0, 0), text, font=face)
        x = (1900 - (bounds[2] - bounds[0])) // 2
        y = (190 - (bounds[3] - bounds[1])) // 2 - bounds[1]
        draw.text((x, y), text, font=face, fill=f"#{BLACK}")
        canvas.save(EQUATIONS / f"{name}.png", dpi=(300, 300))


def create_handbook_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)

    img = Image.new("RGB", (1700, 760), "white")
    draw = ImageDraw.Draw(img)
    draw.text((55, 30), "One learner action and the system work behind it", font=font(36, True), fill=f"#{BLACK}")
    draw_box(draw, (60, 140, 430, 360), "Learner sees", ["One clear Learn Next skill", "Why it was chosen", "Start lesson action"], fill=PALE_BLUE, border=BLUE)
    draw_box(draw, (660, 105, 1050, 395), "Application decides", ["Reads global knowledge", "Checks retention and gates", "Builds eligible candidates", "Calls graph and model", "Persists path and reasons"], fill=PALE_TEAL, border=TEAL)
    draw_box(draw, (1280, 140, 1640, 360), "Reviewer can inspect", ["Evidence and thresholds", "Feature and model versions", "Locked and eligible lanes"], fill=PALE_AMBER, border=AMBER)
    draw_arrow(draw, (435, 250), (655, 250), BLUE)
    draw_arrow(draw, (1055, 250), (1275, 250), AMBER)
    draw_box(draw, (660, 500, 1050, 680), "New assessed evidence", ["Updates global skill state", "Marks affected paths stale", "Requires explicit regeneration"], fill=PALE_GRAY, border=NAVY)
    draw_arrow(draw, (1480, 365), (1055, 545), AMBER)
    draw_arrow(draw, (660, 590), (250, 365), TEAL)
    img.save(ASSETS / "visible-and-technical.png")

    img = Image.new("RGB", (1700, 780), "white")
    draw = ImageDraw.Draw(img)
    draw.text((55, 30), "How an assessment changes the path", font=font(36, True), fill=f"#{BLACK}")
    nodes = [
        ("Submit", ["Server scores", "correct option hidden"]),
        ("Evidence", ["Append immutable", "source and context"]),
        ("Knowledge", ["Mastery confidence", "state retention"]),
        ("Invalidate", ["Relevant ACTIVE", "path becomes STALE"]),
        ("Regenerate", ["Gates graph model", "new ACTIVE version"]),
    ]
    positions = [(45, 150, 325, 350), (375, 150, 655, 350), (705, 150, 985, 350), (1035, 150, 1315, 350), (1365, 150, 1645, 350)]
    colors = [(PALE_BLUE, BLUE), (PALE_GRAY, NAVY), (PALE_TEAL, TEAL), (PALE_AMBER, AMBER), (PALE_BLUE, BLUE)]
    for i, (title, lines) in enumerate(nodes):
        draw_box(draw, positions[i], title, lines, fill=colors[i][0], border=colors[i][1])
        if i < len(nodes) - 1:
            draw_arrow(draw, (positions[i][2] + 5, 250), (positions[i + 1][0] - 5, 250), colors[i][1])
    draw.rounded_rectangle((235, 500, 1465, 690), radius=20, outline=f"#{BORDER}", width=3, fill=f"#{PALE_GRAY}")
    draw.text((275, 525), "History remains available", font=font(25, True), fill=f"#{BLACK}")
    draw.text((275, 575), "The former path is SUPERSEDED. The new path records the new evidence state, model contract and reasons.", font=font(21), fill=f"#{MUTED}")
    draw.text((275, 620), "The reviewer can compare what moved, unlocked, became revision due or disappeared.", font=font(21), fill=f"#{MUTED}")
    img.save(ASSETS / "assessment-to-path.png")

    img = Image.new("RGB", (1700, 780), "white")
    draw = ImageDraw.Draw(img)
    draw.text((55, 30), "Cross course knowledge reuse", font=font(36, True), fill=f"#{BLACK}")
    draw_box(draw, (80, 145, 480, 355), "Course A evidence", ["Recursion assessment", "Mastery 0.82", "Confidence 0.71"], fill=PALE_BLUE, border=BLUE)
    draw_box(draw, (650, 120, 1050, 380), "Global Skill Passport", ["One learner skill row", "recursion mastery 0.82", "evidence remains reusable", "course progress stays separate"], fill=PALE_TEAL, border=TEAL)
    draw_box(draw, (1220, 145, 1620, 355), "Course B decision", ["Recursion prerequisite", "recognized as satisfied", "dependent can unlock"], fill=PALE_AMBER, border=AMBER)
    draw_arrow(draw, (485, 250), (645, 250), BLUE)
    draw_arrow(draw, (1055, 250), (1215, 250), AMBER)
    draw_box(draw, (430, 510, 1270, 700), "What does not transfer", ["Module completion is still owned by each enrollment", "LearnPath reuses knowledge evidence, not navigation progress"], fill=PALE_GRAY, border=NAVY)
    img.save(ASSETS / "cross-course.png")


def cover(doc: Document):
    doc.add_paragraph().paragraph_format.space_after = Pt(52)
    tag = doc.add_paragraph("LEARNPATH")
    tag.paragraph_format.space_after = Pt(12)
    set_run_font(tag.runs[0], size=14, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("Reviewer Explanation Handbook")
    set_run_font(r, name="Aptos Display", size=34, bold=True, color=BLACK)
    subtitle = doc.add_paragraph("Technical and nontechnical explanations for presenting and defending the LearnPath project")
    subtitle.paragraph_format.space_after = Pt(30)
    set_run_font(subtitle.runs[0], name="Aptos Display", size=16, color=BLACK)

    para(doc, "This handbook explains the completed LearnPath system in the language a reviewer can understand and at the technical depth a panel can challenge. Each core module includes a plain explanation, the implemented data and algorithm, a worked example, the user-visible proof, and the limits of the claim.")

    add_table(
        doc,
        ["Use this handbook when", "What to read"],
        [
            ["You need a short project explanation", "The one minute and five minute presentation scripts"],
            ["A reviewer asks how a feature works", "The relevant module chapter and request flow"],
            ["A reviewer asks where the AI is", "The machine learning and Knowledge Graph chapters"],
            ["A reviewer challenges accuracy or trust", "The evidence, uncertainty, assurance and limitations chapters"],
            ["You are preparing the live demo", "The demonstration sequence and profile walkthroughs"],
        ],
        widths=[2.5, 4.2],
        size=9.1,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(36)
    meta = doc.add_paragraph("Prepared for final project review panel")
    set_run_font(meta.runs[0], bold=True, size=11, color=BLACK)
    version = doc.add_paragraph("Companion to the detailed technical project report  |  September 2026")
    set_run_font(version.runs[0], size=9.5, color=MUTED)


def contents(doc: Document):
    heading(doc, "How to use this handbook", 1, new_page=True)
    para(doc, "Start with Sections 1 and 2 until you can explain the problem and complete learner flow without technical terms. Then study Sections 3 through 7 to answer implementation questions. Sections 8 through 11 contain worked examples, trust arguments, demo guidance and a large question bank.")
    add_table(
        doc,
        ["Section", "Purpose"],
        [
            ["1 Presentation scripts", "One minute, five minute and technical presentation versions"],
            ["2 Complete system in plain language", "A reviewer friendly explanation of the adaptive loop"],
            ["3 Technical architecture and data ownership", "Service boundaries, database truth and request flow"],
            ["4 Learner knowledge and evidence", "Mastery, confidence, evidence state and the Skill Passport"],
            ["5 Diagnostic and prerequisite intelligence", "Adaptive assessment and Knowledge Graph v2"],
            ["6 Recommendation and adaptation", "Candidates, Random Forest ranking, paths and regeneration"],
            ["7 Supporting learning modules", "Lessons, practice, assessment, retention, feedback and governance"],
            ["8 Worked numerical examples", "Calculations you can reproduce at the panel"],
            ["9 Trust and validation", "What can be trusted, why, and what is still unproven"],
            ["10 Live demonstration guide", "A clear panel flow using the five demo learners"],
            ["11 Reviewer question bank", "Short answers and deeper technical answers"],
            ["12 Glossary and final revision sheet", "Key terms and the last page summary"],
        ],
        widths=[2.2, 4.5],
    )
    heading(doc, "The explanation pattern", 2)
    add_table(
        doc,
        ["Layer", "Question it answers"],
        [
            ["Plain explanation", "What does this feature do for the learner"],
            ["Technical explanation", "What data, policy and service implement it"],
            ["Worked example", "Can we follow the decision with numbers or a concrete scenario"],
            ["Visible proof", "Where can a reviewer see real backend behavior in the interface"],
            ["Trust boundary", "What is deterministic, what is predicted and what remains uncertain"],
        ],
        widths=[1.8, 4.9],
    )


def presentation_scripts(doc: Document):
    heading(doc, "1 Presentation scripts", 1, new_page=True)
    heading(doc, "One minute explanation", 2)
    panel_words(doc, "LearnPath decides what a learner should study next. It keeps one global record of what the learner knows, so knowledge demonstrated in one course can satisfy a prerequisite in another. When evidence is missing, a short adaptive diagnostic gathers the most useful observations without pretending that one answer proves mastery. The prerequisite graph removes unsafe choices, a Random Forest ranks the remaining choices by predicted learning benefit, and the system creates a path that shows recognized knowledge, the current skill, Learn Next, upcoming skills and locked skills. After practice or assessment changes mastery or retention, the path becomes stale and is regenerated as a new version. The learner can therefore see what they know, what is missing, why a skill is locked and why the next skill was selected.")

    heading(doc, "Five minute explanation", 2)
    numbered(doc, [
        "Begin with the limitation of a fixed course sequence. Two students who enter the same course may have different prior knowledge and different forgotten skills.",
        "Explain the global Skill Passport. A skill belongs to the learner, while course progress belongs to an enrollment. This separation is what makes cross-course reuse possible.",
        "Explain evidence. Diagnostics, practice, assessment and retention checks create server-scored observations. Lesson completion records activity but never creates mastery.",
        "Explain the diagnostic. It has a maximum of 15 questions and selects coverage, confirmation and verification items. A single answer is only a probe. Untested skills remain explicitly untested.",
        "Explain the Knowledge Graph. TypeScript enforces required prerequisite thresholds. NetworkX adds graph structure such as layers, bottlenecks, shortest routes and counterfactual unlocks.",
        "Explain ML. Only eligible candidates go to the 55-feature Random Forest model. The model predicts relative benefit among safe choices; it cannot unlock a skill.",
        "Explain the path. The API stores every lane, score, learner state, reason code and version used at decision time.",
        "Close the loop. New assessed evidence marks the path stale. Regeneration runs the entire pipeline again and preserves the previous version for comparison.",
    ])

    heading(doc, "Technical explanation in ten points", 2)
    add_table(
        doc,
        ["Point", "Technical statement"],
        [
            ["1", "PostgreSQL is the source of truth for curriculum, learner state, evidence and versioned decisions"],
            ["2", "The global learner skill key is learner_id plus skill_id, not learner plus course plus skill"],
            ["3", "The Express TypeScript API owns all authoritative policy and transaction boundaries"],
            ["4", "Evidence updates use source weights, difficulty-adjusted performance and confidence policy"],
            ["5", "Diagnostic v3 selects one question at a time under a 15-item evidence budget"],
            ["6", "Required prerequisite gates use confidence-adjusted mastery before any ranking"],
            ["7", "NetworkX receives a read-only request-scoped graph projection and cannot mutate learner state"],
            ["8", "Random Forest inference verifies checksum, versions and the exact 55-feature order"],
            ["9", "The path transaction stores model, feature, inference and policy provenance"],
            ["10", "Evidence invalidation and explicit regeneration preserve historical path versions"],
        ],
        widths=[0.55, 6.15],
    )
    heading(doc, "Statements to avoid", 2)
    add_table(
        doc,
        ["Avoid saying", "Say this instead"],
        [
            ["The AI decides everything", "Deterministic policy decides safety and learner state; ML ranks safe candidates"],
            ["The diagnostic knows every skill after 15 questions", "It gathers the highest-value evidence and returns untested skills honestly"],
            ["The model is highly accurate", "The frozen synthetic evaluation is reproducible, but real-world effectiveness still requires a learner study"],
            ["Completing a lesson improves mastery", "Only scored performance evidence updates mastery"],
            ["We use Neo4j for the knowledge graph", "PostgreSQL stores graph truth and NetworkX performs derived graph algorithms"],
            ["A recommendation caused the learner improvement", "The system observed later improvement after acceptance; the association is not proof of causation"],
        ],
        widths=[2.3, 4.4],
    )


def plain_system(doc: Document):
    heading(doc, "2 Complete system in plain language", 1, new_page=True)
    para(doc, "A course normally asks every learner to follow the same order. LearnPath instead treats the order as a decision that can change. The decision starts from evidence about the learner and ends with one explainable next action.")
    add_picture(doc, ASSETS / "visible-and-technical.png", 6.65, "The simple learner action is backed by a complete auditable decision", "Learner view, application decision process, reviewer evidence and later adaptation")
    heading(doc, "The eight questions LearnPath answers", 2)
    add_table(
        doc,
        ["Question", "Answer produced by the system"],
        [
            ["What does the learner know", "Global mastery, confidence and evidence state for each skill"],
            ["How certain are we", "Observation count, multiple sessions, source diversity, consistency and confidence"],
            ["What may have been forgotten", "A retention projection anchored to the latest scored evidence"],
            ["What is required for this course", "Course skill targets and recursive required prerequisite ancestors"],
            ["What is safe now", "Eligible candidates after all required gates pass"],
            ["What is the best safe option", "Random Forest benefit probability plus bounded policy adjustments"],
            ["Why was it selected", "A stored explanation using gap, retention, graph leverage, course and goal context"],
            ["What changed later", "A stale signal and a versioned regenerated path with a change summary"],
        ],
        widths=[2.0, 4.7],
    )
    heading(doc, "A simple learner story", 2)
    para(doc, "Assume a learner enters an algorithms course. They already demonstrated recursion in another course, have weak evidence for sorting, and have not been tested on dynamic programming. LearnPath recognizes recursion globally, keeps dynamic programming locked behind the required foundation, and creates sorting as an eligible learning candidate. The model ranks sorting against other eligible skills. The path shows recursion as recognized, sorting as Learn Next and dynamic programming as locked with the exact missing prerequisite.")
    para(doc, "After the learner studies sorting, lesson completion alone does not change mastery. Practice supplies a low-weight observation, and a separate assessment supplies a stronger observation. If the global sorting state changes, the current path becomes stale. The learner chooses regeneration, and the new path may unlock a dependent algorithm skill. The old path remains in history.")
    panel_words(doc, "The system does not force a different order simply because it can. It changes the order only when learner evidence, retention or prerequisite readiness gives a reason.")


def architecture(doc: Document):
    heading(doc, "3 Technical architecture and data ownership", 1, new_page=True)
    add_picture(doc, REPORT_ASSETS / "architecture.png", 6.55, "The API owns decisions and Python contributes bounded intelligence", "Architecture with React, Express, PostgreSQL and FastAPI boundaries")
    heading(doc, "Why four layers", 2)
    add_table(
        doc,
        ["Layer", "What it owns", "Why this boundary is important"],
        [
            ["React and Vite web client", "Navigation, visual explanation, forms and learner actions", "The browser cannot invent scores or bypass a prerequisite"],
            ["Express TypeScript API", "Authentication, validation, mastery, diagnostics, retention, prerequisites, candidates and paths", "All authoritative rules execute in one application boundary"],
            ["PostgreSQL", "Curriculum, global skill state, evidence, enrollments, path history and feedback", "Transactions and constraints protect a single source of truth"],
            ["FastAPI intelligence service", "Random Forest inference and NetworkX graph analytics", "Stateless advisory work can fail safely without corrupting learner state"],
        ],
        widths=[1.55, 2.75, 2.4],
    )
    heading(doc, "A path generation request", 2)
    numbered(doc, [
        "The browser sends an authenticated generate request for one enrollment.",
        "The API confirms that the enrollment belongs to the authenticated learner.",
        "The API loads course skills, recursive prerequisite ancestors, global learner state, recent activity and retention.",
        "The TypeScript prerequisite engine decides which required edges pass and which candidates are locked.",
        "The API optionally sends a read-only graph projection to NetworkX for structural metrics.",
        "The candidate engine creates eligible, locked and excluded lanes. It asserts that only eligible candidates continue.",
        "The API builds the exact 55-field feature record for each eligible candidate and calls FastAPI inference.",
        "FastAPI verifies the artifact checksum and feature contract, then returns benefit probabilities.",
        "The API applies bounded revision and graph bonuses, orders dependency-safely and writes a path transaction.",
        "The browser receives the persisted path and renders its lanes and explanations.",
    ])
    heading(doc, "Single source of truth", 2)
    para(doc, "LearnPath deliberately avoids storing the same mastery in an enrollment table, goal table and learner profile. The only knowledge record is identified by learner_id and skill_id. Enrollment owns course progress. A goal adds relevance but does not own knowledge. This decision prevents three values for the same skill from disagreeing.")
    add_table(
        doc,
        ["Fact", "Owner"],
        [
            ["The learner scored 8 out of 10 on a recursion assessment", "Append-only skill evidence"],
            ["The current recursion mastery estimate is 0.82", "Global learner skill state"],
            ["The learner completed Module 3 in Course A", "Enrollment module progress"],
            ["Recursion is relevant to the interview goal", "Learner goal mapping"],
            ["Course B recognizes recursion as a prerequisite", "Live prerequisite projection from global state"],
        ],
        widths=[4.45, 2.25],
    )
    code_location(doc, ["apps/web/src", "apps/api/src/app.ts", "apps/api/src", "database/migrations", "services/ml/app"])


def knowledge_evidence(doc: Document):
    heading(doc, "4 Learner knowledge and evidence", 1, new_page=True)
    heading(doc, "Global Skill Passport", 2)
    para(doc, "Plain explanation  The Skill Passport is the learner's reusable knowledge record. If the learner proves a skill in one course, every other course that uses the same canonical skill can read that knowledge. The learner does not have to repeat an already demonstrated prerequisite merely because they enrolled in another course.", lead="Plain explanation  ")
    para(doc, "Technical explanation  PostgreSQL enforces one row per learner_id and skill_id. The learner-skills service joins that row to the active skill catalog, evidence counts, retention and course mappings. It returns the same mastery and confidence in every course context while leaving course progress inside the enrollment.", lead="Technical explanation  ")
    add_picture(doc, ASSETS / "cross-course.png", 6.55, "Knowledge transfers through the global skill key while course completion stays local", "Course A evidence updates a global recursion skill that Course B can reuse")
    panel_words(doc, "We reuse demonstrated knowledge, not course completion. That distinction prevents a learner from receiving mastery merely by opening or completing content.")
    code_location(doc, ["apps/api/src/learner-skills/service.ts", "database/migrations/0006_global_learner_skill_model.sql"])

    heading(doc, "Evidence ledger", 2)
    para(doc, "Plain explanation  Every mastery value needs a reason. LearnPath therefore keeps the performance observations that produced it. A reviewer can inspect whether the evidence came from a diagnostic, practice, assessment or retention check and see the before and after state.", lead="Plain explanation  ")
    para(doc, "Technical explanation  skill_evidence is append-only. Database triggers reject update and delete operations. Each row preserves learner, skill, source type, score, correctness, difficulty, session and attempt information, timestamps, state before and after, retention values and structured metadata. The application commits the evidence row and learner-skill update atomically.", lead="Technical explanation  ")
    add_table(
        doc,
        ["Evidence source", "Weight", "Why the weight differs"],
        [
            ["Diagnostic", "0.40 multiplied by diagnostic reliability", "Placement is useful but one session remains bounded"],
            ["Practice", "0.25", "Practice is formative and may include trial and error"],
            ["Quiz", "0.30", "More structured than practice but usually narrower than assessment"],
            ["Assessment", "0.40", "Separate scored assessment is stronger direct evidence"],
            ["Module assessment", "0.40", "Summative performance evidence"],
            ["Retention check", "0.35", "Directly tests whether earlier knowledge remains accessible"],
        ],
        widths=[1.6, 1.0, 4.1],
    )
    panel_words(doc, "The weights control how fast an estimate moves. They do not convert activity into knowledge; only server-scored performance reaches this engine.")
    code_location(doc, ["apps/api/src/skill-evidence/service.ts", "apps/api/src/mastery/service.ts", "apps/api/src/mastery/policy.ts"])

    heading(doc, "Mastery update", 2)
    para(doc, "Plain explanation  Mastery is a moving estimate. New evidence pulls the old estimate toward the measured performance. Stronger evidence pulls more; weaker evidence pulls less. The result stays between zero and one.", lead="Plain explanation  ")
    equation(doc, "mastery", 5.9, "Mastery update used for scored evidence", "M is mastery, P is measured performance, wsource is the configured source weight, r is reliability and clip keeps the result from zero to one.")
    para(doc, "Technical detail  When prior mastery exists, the source weight comes from policy. A new skill begins from a neutral 0.50 prior. Diagnostic evidence uses the same engine but multiplies the source weight by reliability. Difficulty-adjusted performance gives more credit for a correct hard answer and a stronger negative signal for an incorrect easy answer.")
    panel_words(doc, "We use incremental updating rather than replacing mastery with the latest score. One unusual attempt therefore cannot erase or create a complete knowledge profile.")

    heading(doc, "Confidence and evidence state", 2)
    para(doc, "Plain explanation  Mastery answers how well the learner appears to know the skill. Confidence answers how much evidence supports that estimate. A learner can have high estimated mastery with low confidence if the estimate comes from only one or two observations.", lead="Plain explanation  ")
    equation(doc, "confidence", 6.2, "Confidence grows from repeated and varied evidence", "n is observation count, D is source diversity, S is session count, d is average difficulty, K is consistency and R is one when retention evidence exists. Diagnostic-only confidence is capped at 0.45.")
    add_table(
        doc,
        ["Evidence state", "Rule", "How to explain it"],
        [
            ["UNKNOWN", "No scored observation", "We do not know yet"],
            ["ESTIMATED", "Some evidence but below assessed thresholds", "Useful for personalization but still uncertain"],
            ["ASSESSED", "At least 5 observations, 2 sessions and confidence 0.45", "Repeated assessment evidence supports the estimate"],
            ["VERIFIED", "At least 12 observations, 4 sessions, consistency 0.70, confidence 0.78 and evidence diversity", "The strictest implemented evidence standard"],
        ],
        widths=[1.2, 3.2, 2.3],
    )
    panel_words(doc, "A high mastery number alone is never enough to call a skill verified. Verification requires evidence volume, multiple sessions, consistency, confidence and diversity.")


def diagnostic_graph(doc: Document):
    heading(doc, "5 Diagnostic and prerequisite intelligence", 1, new_page=True)
    heading(doc, "Adaptive diagnostic v3", 2)
    para(doc, "Plain explanation  The diagnostic is a short placement conversation. It first samples different important skills. It then returns to uncertain or contradictory skills. The final questions try to verify application-level knowledge. This gives a useful starting point without pretending to examine every skill exhaustively.", lead="Plain explanation  ")
    add_picture(doc, REPORT_ASSETS / "diagnostic.png", 6.5, "The diagnostic spends its fixed budget on coverage confirmation and verification", "Adaptive diagnostic stages and the rule that one answer cannot establish mastery")
    para(doc, "Technical explanation  The server reveals exactly one question at a time. It selects from active skill-linked questions in the course. The fixed maximum is 15 questions: five coverage slots, eight confirmation slots and two verification slots. Correctness stays hidden until final submission, so an early item cannot coach later answers. Answer drafts allow exact resume.", lead="Technical explanation  ")
    heading(doc, "How the selector scores a question", 3)
    add_table(
        doc,
        ["Signal", "Implemented effect"],
        [
            ["Unknown skill", "Adds 100 to base skill priority"],
            ["Mastery gap", "Adds 50 times the distance from course target"],
            ["Uncertainty", "Adds 20 times one minus current confidence"],
            ["Graph gateway", "Adds 4 for each required downstream dependent"],
            ["Retention risk", "Adds up to 25 for critical or at-risk knowledge"],
            ["Discrimination", "Adds 8 times the item's bounded discrimination"],
            ["Cognitive novelty", "Adds 12 when the cognitive level is new for that skill"],
            ["Application needed", "Adds 18 when APPLY or ANALYZE evidence is still missing"],
            ["Difficulty fit", "Favors difficulty close to the level implied by current mastery"],
            ["Confirmation need", "Strongly favors the second or third observation for unresolved evidence"],
        ],
        widths=[1.65, 5.05],
    )
    para(doc, "The selector is a transparent deterministic policy, not an item-response-theory model. Its job is to spend a small question budget on useful evidence. The item metadata includes cognitive level, discrimination, guessing prior, calibration state and optional misconception codes so later real response data can support formal calibration.")
    heading(doc, "How answers become a cautious result", 3)
    para(doc, "A correct item produces a difficulty-aware measurement value of 0.72 plus 0.05 times difficulty. An unsure incorrect response produces 0.12 plus 0.035 times difficulty. Another incorrect response produces 0.03 plus 0.05 times difficulty. The skill performance score is the discrimination-weighted average of these values, where discrimination is clipped between 0.25 and 2.5.")
    equation(doc, "diagnostic_reliability", 5.3, "Diagnostic reliability is bounded even when all answers agree", "n is direct observations, L is distinct cognitive levels, A is one when application evidence is present, and K is one when all answers are consistent. The maximum reliability is 0.78.")
    equation(doc, "mastery_interval", 5.8, "The result reports an uncertainty interval around mastery", "M is mastery, C is confidence and n is direct observations. More observations and higher confidence narrow the interval.")
    add_table(
        doc,
        ["Classification", "Implemented meaning"],
        [
            ["PROBED", "Exactly one direct observation; never mastery certification"],
            ["READY", "At least two correct observations and mastery at least 0.55, but below mastery certification"],
            ["MASTERED", "At least three observations, a correct APPLY or ANALYZE item, target mastery and confidence at least 0.35"],
            ["GAP", "At least two observations, no correct answers and mastery below 0.45"],
            ["NEEDS CONFIRMATION", "Mixed or insufficient evidence"],
            ["FORGOTTEN", "Earlier mastery met target, retention is at risk or critical, and current performance is below 0.45"],
            ["FRAGILE FOUNDATION", "Advanced knowledge coexists with a weak tested required prerequisite"],
            ["NOT TESTED", "No direct item was selected inside the fixed budget"],
        ],
        widths=[1.65, 5.05],
    )
    panel_words(doc, "Fifteen questions do not certify all course skills. The engine prioritizes the most informative evidence and keeps untested skills visible. Later practice and assessment continue the estimation process.")
    code_location(doc, ["apps/api/src/diagnostics/selection.ts", "apps/api/src/diagnostics/estimation.ts", "apps/api/src/diagnostics/service.ts"])

    heading(doc, "Knowledge Graph v2", 2, new_page=True)
    para(doc, "Plain explanation  The graph describes which skills depend on which foundations. It stops LearnPath from recommending an advanced skill before the learner is ready and explains the sequence needed to reach it.", lead="Plain explanation  ")
    add_picture(doc, REPORT_ASSETS / "knowledge-graph.png", 5.9, "Graph truth safety analytics and ranking remain separate", "Knowledge Graph v2 with PostgreSQL, TypeScript, NetworkX, Random Forest and a safe path")
    para(doc, "Technical explanation  PostgreSQL stores canonical skills, course mappings and typed prerequisite edges. TypeScript recursively loads the relevant subgraph and makes every REQUIRED gate decision. NetworkX analyzes a read-only projection. The graph is a directed acyclic graph, so topological layers and dependency order are defined.", lead="Technical explanation  ")
    heading(doc, "Required and recommended edges", 3)
    para(doc, "A REQUIRED edge blocks the dependent skill until the prerequisite reaches the edge threshold. A RECOMMENDED edge adds guidance but does not block. Unknown mastery never satisfies a required edge. Low confidence reduces the mastery value used by the gate.")
    equation(doc, "prerequisite", 5.4, "Confidence adjusted prerequisite gate", "C is confidence, M is observed mastery and T is the threshold stored on the prerequisite edge. The largest possible uncertainty penalty is 0.10.")
    panel_words(doc, "Observed mastery 0.73 with confidence 0.40 becomes effective mastery 0.67. If the edge requires 0.70, the skill stays locked. This is deliberate caution, not a model prediction.")
    heading(doc, "What NetworkX adds", 3)
    add_table(
        doc,
        ["Algorithm", "What it tells us", "Use in LearnPath"],
        [
            ["Topological generations", "Which skills belong to each dependency level", "Learner-facing graph layers and valid order"],
            ["Descendants and direct dependents", "How much later learning relies on one foundation", "Downstream reach"],
            ["Betweenness centrality", "Whether a skill connects graph regions", "Structural bridge importance"],
            ["Shortest path", "A minimal prerequisite route from a foundation to a target", "Explainable foundation route"],
            ["Bottleneck analysis", "Which weak skills currently block many descendants", "Focused prerequisite attention"],
            ["Counterfactual unlocks", "What would become ready if one selected prerequisite met target", "Explain the value of closing a specific gap"],
        ],
        widths=[1.55, 2.65, 2.5],
    )
    equation(doc, "gateway", 3.6, "NetworkX gateway score", "Dr is descendant ratio, Bn is normalized betweenness and Ur is immediate counterfactual unlock ratio. Every input is bounded from zero to one.")
    heading(doc, "Why NetworkX and not Neo4j", 3)
    para(doc, "NetworkX is used because the project needs in-memory algorithms on a small request-scoped graph. PostgreSQL already stores the graph together with curriculum constraints and learner state. Adding Neo4j would create a second authoritative copy that must be synchronized. At the implemented scale of five courses and 36 shared skills, that operational cost is not justified. Neo4j becomes attractive when graph-native queries, heterogeneous relationships, very large traversals or independent graph operations dominate the workload.")
    panel_words(doc, "We are using graph algorithms. We are not using a graph database because the relational curriculum is already the safest source of truth at this scale.")
    code_location(doc, ["apps/api/src/prerequisites/engine.ts", "apps/api/src/prerequisites/service.ts", "services/ml/app/graph.py", "docs/knowledge-graph-v2.md"])


def recommendation_adaptation(doc: Document):
    heading(doc, "6 Recommendation and adaptation", 1, new_page=True)
    heading(doc, "Candidate generation", 2)
    para(doc, "Plain explanation  Before asking the model which skill is best, LearnPath creates a safe list. It separates skills that can be learned now, skills that are locked and skills that are already satisfied.", lead="Plain explanation  ")
    para(doc, "Technical explanation  The candidate engine combines the enrollment's course skills with recursive required ancestors. It joins global mastery, confidence, evidence state, retention, module context, goal relevance, content availability and prerequisite results. Each candidate is LEARN, REVISION or SUPPORTING_PREREQUISITE and is placed in ELIGIBLE, LOCKED or EXCLUDED.", lead="Technical explanation  ")
    add_table(
        doc,
        ["Lane", "Rule", "What happens next"],
        [
            ["ELIGIBLE", "Target is not already satisfied and no required prerequisite is missing", "Feature construction and ML inference"],
            ["LOCKED", "At least one required prerequisite gate fails", "Displayed with exact missing prerequisites; never sent to ML"],
            ["EXCLUDED", "Global mastery meets course target and no revision signal is active", "Displayed as recognized prior knowledge"],
        ],
        widths=[1.1, 3.7, 1.9],
    )
    panel_words(doc, "Candidate generation is where we protect the model boundary. A high model score can never rescue a locked skill because the model never receives that row.")
    code_location(doc, ["apps/api/src/candidates/engine.ts", "apps/api/src/candidates/service.ts"])

    heading(doc, "Machine learning ranking", 2)
    para(doc, "Plain explanation  The model compares the safe options and predicts which one is most likely to produce useful learning. It considers current knowledge, recent performance, evidence strength, retention, prerequisite readiness and engagement history.", lead="Plain explanation  ")
    para(doc, "Technical explanation  The deployed contract is learner-candidate-features-v2 with 55 numeric features in fixed order. Identifiers and the label remain outside the matrix. Current outcomes such as post-assessment, mastery gain, feedback and completion are blocked. Historical outcome features are shifted by one interaction, so the model cannot see the answer it is supposed to predict.", lead="Technical explanation  ")
    heading(doc, "Training data and experiment", 3)
    add_table(
        doc,
        ["Element", "Implemented value"],
        [
            ["Data", "1,000 synthetic learners and 20,000 interactions, seed 42"],
            ["Label", "Benefit score at least 0.27 using mastery gain, completion, assessment improvement and feedback"],
            ["Split", "700 train learners, 150 validation learners and 150 test learners with no learner overlap"],
            ["Models", "Gradient Boosting, Random Forest and Logistic Regression"],
            ["Baselines", "Highest Skill Gap and curriculum-derived Popularity"],
            ["Selection", "Validation NDCG at 5, then ROC-AUC and F1 tie-breaks"],
            ["Selected model", "Random Forest with 220 trees, depth 12, minimum leaf 4 and balanced class weights"],
        ],
        widths=[1.55, 5.15],
    )
    add_table(
        doc,
        ["Candidate", "Validation NDCG at 5", "Test NDCG at 5", "Test ROC AUC", "Test Precision at 5"],
        [
            ["Random Forest", "0.581021", "0.578463", "0.801410", "0.525333"],
            ["Gradient Boosting", "0.576017", "0.589501", "0.804062", "0.521333"],
            ["Logistic Regression", "0.564225", "0.580259", "0.802338", "0.513333"],
            ["Highest Skill Gap", "0.576408", "0.581229", "0.730785", "0.510667"],
            ["Popularity", "0.251677", "0.272412", "0.440132", "0.284000"],
        ],
        widths=[1.65, 1.3, 1.2, 1.2, 1.35],
        size=8.2,
    )
    para(doc, "How to interpret the result  Random Forest won by the declared validation selection rule. Its test ROC-AUC and Precision at 5 exceed the Highest Skill Gap baseline, but its test NDCG at 5 is slightly lower. This mixed result supports using the model as a bounded ranking component. It does not support claiming that the synthetic model is universally better or educationally proven.", lead="How to interpret the result  ")
    heading(doc, "Checked inference", 3)
    bullets(doc, [
        "FastAPI verifies the artifact SHA-256 before deserialization.",
        "The model, feature and inference versions must match the manifest.",
        "The request must contain all 55 feature names in exact order and valid bounds.",
        "A bad request returns 422. A missing or incompatible artifact returns 503.",
        "There is no heuristic probability fallback and the service never writes PostgreSQL.",
    ])
    panel_words(doc, "We can trust the serving contract to reject the wrong artifact or schema. We cannot yet claim real learner effectiveness because the training evaluation is synthetic.")
    code_location(doc, ["services/ml/app/inference.py", "services/ml/models/benefit-ranking-v2/manifest.json", "apps/api/src/paths/features.ts", "apps/api/src/paths/ml-client.ts"])

    heading(doc, "Personalized path", 2, new_page=True)
    para(doc, "Plain explanation  The path shows what the learner already knows, what they are doing, the next recommended skill, later safe work and skills that remain locked. It is the final explanation of the system's reasoning.", lead="Plain explanation  ")
    para(doc, "Technical explanation  Every eligible candidate has a Random Forest benefit probability. The API adds a fixed revision bonus when a previously mastered skill is due for review and a bounded graph bonus based on gateway value. It sorts by total priority, probability, dependency level, course sequence and skill name. Locked and recognized skills are inserted as separate lanes without model probabilities.", lead="Technical explanation  ")
    equation(doc, "priority", 4.5, "Final priority among eligible candidates", "The probability comes from Random Forest. The revision indicator adds 0.08. G is the zero-to-one NetworkX gateway score, so the largest graph contribution is 0.06.")
    add_table(
        doc,
        ["Path lane", "Meaning"],
        [
            ["RECOGNIZED", "Global mastery already meets this course target"],
            ["CURRENT", "Recently active in-progress learning stays visible"],
            ["RECOMMENDED NEXT", "Highest priority eligible candidate"],
            ["UPCOMING", "Other eligible candidates in dependency-aware order"],
            ["LOCKED", "Required prerequisite is missing and candidate was excluded from ML"],
        ],
        widths=[1.85, 4.85],
    )
    para(doc, "The stored path snapshot includes mastery, confidence, evidence state, retention, prerequisite checks, model probability, policy bonuses, goal relevance, course context, reason codes, explanations and the model, feature, inference and policy versions. A reviewer can therefore reconstruct why the item appeared where it did.")
    panel_words(doc, "The model suggests priority only among safe candidates. The path engine still applies dependency order, bounded policy bonuses and deterministic tie-breaks.")
    code_location(doc, ["apps/api/src/paths/engine.ts", "apps/api/src/paths/policy.ts", "apps/api/src/paths/service.ts"])

    heading(doc, "Cross course coordination", 2)
    para(doc, "If two active course paths recommend the same global skill, the coordination service can prefer that skill because one learning action helps more than one course. It groups current ACTIVE paths by skill, keeps the strongest course priority, adds up to two shared-course bonuses and optionally adds goal relevance.")
    equation(doc, "coordination", 4.4, "Cross course coordination score", "Qbest is the strongest course path priority, k is the number of active course contexts for that skill and g is goal relevance from zero to one.")
    panel_words(doc, "Coordination does not merge the courses. It selects one learner-facing next action while preserving the independent dependency-safe path for every enrollment.")
    code_location(doc, ["apps/api/src/paths/coordination.ts", "apps/api/src/paths/policy.ts"])

    heading(doc, "Dynamic path regeneration", 2)
    add_picture(doc, ASSETS / "assessment-to-path.png", 6.55, "Assessed evidence invalidates a path before a new version is generated", "Assessment submission, immutable evidence, knowledge update, stale path and explicit regeneration")
    para(doc, "Plain explanation  LearnPath does not silently rewrite the learner's path. When new evidence matters, it marks the current path as stale and asks for regeneration. The learner and reviewer can compare the former and new versions.", lead="Plain explanation  ")
    para(doc, "Technical explanation  The lifecycle is ACTIVE to STALE to SUPERSEDED, followed by a newly created ACTIVE path. Regeneration locks the current path row, reruns prerequisite analysis, NetworkX, candidate generation, feature construction, inference and ordering, then inserts the new path, items, events and change summary in one transaction.", lead="Technical explanation  ")
    panel_words(doc, "The old recommendation remains evidence of what the system knew at that time. This is why path versioning matters for trust and evaluation.")
    code_location(doc, ["apps/api/src/paths/change.ts", "apps/api/src/paths/service.ts", "database/migrations/0020_dynamic_path_regeneration.sql"])


def supporting_modules(doc: Document):
    heading(doc, "7 Supporting learning modules", 1, new_page=True)
    heading(doc, "Learning resources practice and assessment", 2)
    add_table(
        doc,
        ["Activity", "What it records", "Does it update mastery"],
        [
            ["Open or complete lesson", "Activity event, resource status and time", "No"],
            ["Practice answer", "Server-scored PRACTICE evidence", "Yes with weight 0.25"],
            ["Quiz answer", "Server-scored QUIZ evidence", "Yes with weight 0.30"],
            ["Separate assessment", "Server-scored ASSESSMENT evidence", "Yes with weight 0.40"],
            ["Retention check", "Server-scored RETENTION_CHECK evidence", "Yes with weight 0.35"],
        ],
        widths=[1.7, 3.15, 1.85],
    )
    para(doc, "The server owns question selection and correctness. The browser never receives correct-option metadata before submission. Practice confirms that the learner owns the optional enrollment context and that the skill belongs to it. The assessment uses a separate content pool so a learner does not merely repeat the same item they practiced.")
    panel_words(doc, "The visible learning flow and the knowledge engine are connected through scored evidence, not through a fake completion-to-mastery shortcut.")
    code_location(doc, ["apps/api/src/learning/service.ts", "apps/api/src/practice/service.ts", "apps/web/src/pages/StudyResourcePage.tsx"])

    heading(doc, "Retention and forgetting", 2)
    para(doc, "Plain explanation  A learner may have demonstrated a skill earlier but may no longer recall it easily. LearnPath keeps the original mastery estimate and separately predicts current retention. This allows revision without pretending the learner never learned the skill.", lead="Plain explanation  ")
    equation(doc, "retention", 6.35, "Retention projection with confidence and repetition adjusted decay", "C is confidence, n is evidence count, M is mastery and t is days since the latest scored evidence. The base decay value is 0.025.")
    add_table(
        doc,
        ["Retention state", "Range", "Interpretation"],
        [
            ["STRONG", "0.75 to 1.00", "Knowledge is projected to remain readily accessible"],
            ["MODERATE", "0.55 to below 0.75", "Usable knowledge with a future check approaching"],
            ["AT RISK", "0.30 to below 0.55", "Review may be due if mastery was previously demonstrated"],
            ["CRITICAL", "Below 0.30", "Urgent review may be due if mastery was previously demonstrated"],
            ["UNKNOWN", "No evidence anchor", "The system cannot project retention"],
        ],
        widths=[1.25, 1.55, 3.9],
    )
    para(doc, "Revision eligibility is intentionally narrow. Mastery must first have reached at least 0.65, and the retention state must be AT_RISK or CRITICAL. A weak unlearned skill is therefore a learning gap, not a revision recommendation.")
    panel_words(doc, "Mastery stores demonstrated knowledge. Retention estimates current accessibility. Separating them lets us distinguish never learned from learned but forgotten.")
    code_location(doc, ["apps/api/src/retention/service.ts", "apps/api/src/retention/policy.ts"])

    heading(doc, "Recommendation feedback and outcome attribution", 2)
    para(doc, "The learner may accept or decline only the current Learn Next item. A decline can include a structured reason. For an accepted item, later assessed evidence within 30 days is compared with the decision-time baseline. A gain of at least 0.05 becomes IMPROVED, a loss of at least 0.05 becomes DECLINED, a smaller change becomes STABLE, and evidence without a baseline becomes OBSERVED_NO_BASELINE.")
    para(doc, "This process is observational. It associates a later outcome with a versioned recommendation and does not claim that the recommendation caused the outcome. A causal claim would require a controlled study or another defensible experimental design.")
    code_location(doc, ["apps/api/src/recommendations/policy.ts", "apps/api/src/recommendations/service.ts"])

    heading(doc, "Governance", 2)
    add_table(
        doc,
        ["Metric or action", "Minimum sample", "Behavior below minimum"],
        [
            ["Acceptance rates", "30 responses", "Rate is withheld"],
            ["Calibration", "30 assessed outcomes", "Calibration conclusion is withheld"],
            ["Prediction drift", "50 recent and 50 reference predictions", "Drift conclusion is withheld"],
            ["Retraining eligibility", "100 assessed recommendation outcomes", "Retraining remains ineligible"],
        ],
        widths=[2.25, 1.6, 2.85],
    )
    para(doc, "No automatic retraining or promotion occurs. A human must review evidence and deliberately promote a new artifact or policy version. This prevents a small demonstration dataset from changing production behavior automatically.")
    panel_words(doc, "A blank or withheld governance metric is a safety feature when the sample is too small. Showing a percentage would imply evidence that does not exist.")
    code_location(doc, ["apps/api/src/governance/service.ts", "apps/api/src/governance/router.ts"])


def numerical_examples(doc: Document):
    heading(doc, "8 Worked numerical examples", 1, new_page=True)
    heading(doc, "Assessment mastery update", 2)
    para(doc, "Assume current mastery is 0.50 and a separate assessment produces performance 0.90. Assessment weight is 0.40 and reliability is 1.00. The effective weight is 0.40. New mastery is 0.50 times 0.60 plus 0.90 times 0.40, which equals 0.66. The strong assessment moves the estimate by 0.16; it does not replace the entire history with 0.90.")
    add_table(doc, ["Input", "Value"], [["Previous mastery", "0.50"], ["Measured performance", "0.90"], ["Source weight", "0.40"], ["New mastery", "0.66"]], widths=[3.2, 3.5])

    heading(doc, "Diagnostic mastery update", 2)
    para(doc, "Use the same prior mastery 0.50 and measured diagnostic performance 0.90, but assume diagnostic reliability is 0.60. Effective weight becomes 0.40 times 0.60, or 0.24. New mastery is 0.50 times 0.76 plus 0.90 times 0.24, which equals 0.596. The diagnostic moves the estimate less because the bounded session is not certification.")

    heading(doc, "Confidence calculation", 2)
    para(doc, "Assume six observations from two source types over two sessions, average difficulty 3, consistency 0.80 and no retention evidence. The signals are approximately 0.284 from count, 0.10 from diversity, 0.06 from sessions, 0.036 from difficulty and 0.096 from consistency, plus the 0.08 base. Total confidence is approximately 0.656. If every source were diagnostic, the diagnostic-only ceiling would reduce it to 0.45.")

    heading(doc, "Retention after thirty days", 2)
    para(doc, "Assume mastery 0.80, confidence 0.70, ten observations and 30 days since evidence. The confidence factor is 0.80. The repetition factor is about 0.777. Effective decay is about 0.0155. Retention is 0.80 times e raised to minus 0.0155 times 30, which is approximately 0.50. That falls in AT_RISK. Because earlier mastery exceeds 0.65, revision is due.")

    heading(doc, "Prerequisite gate", 2)
    para(doc, "Assume observed mastery 0.73, confidence 0.40 and required threshold 0.70. Uncertainty penalty is 0.06. Effective mastery is 0.67. The required edge fails, so the dependent skill is locked and excluded from ML. If confidence later grows to 0.80 while mastery remains 0.73, the penalty becomes 0.02 and effective mastery becomes 0.71, so the same edge passes.")

    heading(doc, "Path priority", 2)
    para(doc, "Assume Random Forest predicts benefit 0.72. The candidate is due for revision, adding 0.08. NetworkX gateway score is 0.50, adding 0.03 because the graph multiplier is 0.06. Final priority is 0.83. If the skill were locked, this calculation would never occur because it would not enter inference.")

    heading(doc, "Cross course coordination", 2)
    para(doc, "Assume a skill has best course priority 0.78, appears as Learn Next in two active courses and has goal relevance 0.50. One additional course adds 0.04. Goal relevance adds 0.03. Coordinated score is 0.85. The service may present this as the overall next action because one skill advances both course contexts.")

    heading(doc, "Advanced skill with a weak basic foundation", 2)
    para(doc, "Suppose the learner answers an advanced recursion application correctly but misses two basic function-scope questions. LearnPath preserves the recursion evidence rather than forcing a monotonic ability assumption. The weak required foundation remains visible. Diagnostic classification can identify a fragile foundation, the graph shows which advanced skills depend on it, and the path can recommend the foundation while recognizing the advanced evidence. This rare case is exactly why the project stores evidence by skill instead of assigning one overall learner level.")


def trust_validation(doc: Document):
    heading(doc, "9 Trust and validation", 1, new_page=True)
    heading(doc, "Two meanings of trust", 2)
    add_table(
        doc,
        ["Trust question", "Current answer"],
        [
            ["Does the software apply the implemented rules correctly", "Supported by constraints, unit tests, integration tests, version checks and reproducible flows"],
            ["Does the model improve real learner outcomes", "Not yet proven; the model was trained and evaluated on explicitly synthetic data"],
        ],
        widths=[3.1, 3.6],
    )
    para(doc, "The first question is engineering correctness. The second is educational effectiveness. LearnPath provides substantial evidence for the first and clearly limits the second. This distinction is the most important trust statement to make at the panel.")
    heading(doc, "Assurance chain", 2)
    add_table(
        doc,
        ["Risk", "Control", "What the reviewer can inspect"],
        [
            ["One answer creates mastery", "One answer is PROBED; mastery requires repeated varied application evidence", "Diagnostic result and direct observation count"],
            ["Unknown passes a gate", "Unknown never satisfies REQUIRED edges", "Prerequisite explanation and evidence state"],
            ["ML chooses a locked skill", "Locked rows are removed before features and inference", "Candidate partitions and path reason codes"],
            ["Lesson completion creates knowledge", "Resource activity never invokes mastery", "Evidence source history"],
            ["Evidence is edited", "Database rejects update and delete on skill_evidence", "Append-only ledger"],
            ["Wrong model is loaded", "SHA-256, versions, feature order and class contract are checked", "Model manifest and readiness"],
            ["Graph analytics fails", "TypeScript gates continue in safe mode and omit the bonus", "Prerequisite response and service health"],
            ["Future outcomes leak into training", "Blocked fields, shifted histories and learner-disjoint split", "Feature and experiment manifests"],
            ["Path silently changes", "Evidence creates STALE and explicit regeneration creates a new version", "Path history and change summary"],
            ["Small samples create misleading rates", "Governance withholds below minimum counts", "Governance overview"],
        ],
        widths=[1.65, 3.15, 1.9],
        size=8.2,
    )
    heading(doc, "Verification evidence", 2)
    add_table(
        doc,
        ["Area", "Recorded result"],
        [
            ["Latest API tests after Diagnostic v3", "101 passing"],
            ["Latest web tests after Diagnostic v3", "41 passing"],
            ["Documented PostgreSQL integration tests", "17 passing in the final completion audit"],
            ["Documented ML tests", "35 passing in the final completion audit"],
            ["Content coverage", "36 of 36 skills have lesson, diagnostic, practice and separate assessment content"],
            ["Browser flows", "Five learner stories and the complete Aisha adaptation flow passed in the local demo audit"],
        ],
        widths=[3.5, 3.2],
    )
    para(doc, "The integration and ML counts above come from the documented final completion audit. The later Diagnostic v3 work increased the TypeScript totals to 101 API and 41 web tests. The handbook does not present the older integration or ML counts as newly rerun during document generation.")
    heading(doc, "Limitations to state openly", 2)
    bullets(doc, [
        "The Random Forest was trained on synthetic interactions and requires prospective validation with real learners.",
        "Five courses and 36 shared skills demonstrate the architecture but do not validate every learning domain.",
        "Diagnostic item difficulty, discrimination and guessing behavior need calibration from real responses.",
        "Mastery, retention and prerequisite constants are versioned engineering policies that require empirical calibration.",
        "Recommendation attribution is observational and does not establish causality.",
        "Governance correctly withholds rate, calibration and drift conclusions at demonstration sample sizes.",
    ])
    panel_words(doc, "Our claim is that the system is implemented, auditable and dependency-safe. Our next research claim must come from a real learner study, not from the synthetic experiment alone.")


def demo_guide(doc: Document):
    heading(doc, "10 Live demonstration guide", 1, new_page=True)
    heading(doc, "Recommended sequence", 2)
    numbered(doc, [
        "Open Maya Explorer. Show the honest cold start: no enrollment, unknown skill knowledge and no invented path.",
        "Open Noah Starter. Show that enrollment exists but personalization waits for diagnostic evidence.",
        "Open Aisha Builder. Explain the current Learn Next item and its mastery, prerequisite and model reasoning.",
        "Start the lesson. State clearly that resource completion records activity and does not update mastery.",
        "Complete practice and then the separate assessment. Show the server-returned before and after mastery and confidence.",
        "Return to the path. Show that it is STALE, then regenerate and compare the SUPERSEDED and ACTIVE versions.",
        "Open Elena Navigator. Show retention risk, more than one course and coordination across current paths.",
        "Open Ravi Strategist. Show advanced evidence, 20 recognized foundations and the focused Basic Sorting gap.",
        "Open the prerequisite graph. Select a locked node and explain its observed mastery, confidence penalty, target and counterfactual unlocks.",
        "Finish in the reviewer model and governance views. Disclose synthetic training, mixed test results and sample-gated monitoring.",
    ])
    heading(doc, "What each profile proves", 2)
    add_table(
        doc,
        ["Profile", "Technical state", "Reviewer lesson"],
        [
            ["Maya Explorer", "No enrollment and UNKNOWN mastery", "The system does not manufacture knowledge"],
            ["Noah Starter", "Enrolled with incomplete diagnostic", "Evidence precedes personalization"],
            ["Aisha Builder", "Recommendation, learning, assessment, 40-point observed gain and regeneration", "The full adaptive loop works visibly"],
            ["Elena Navigator", "Multiple courses, stale paths and retention review", "Coordination and forgetting are separate from course progress"],
            ["Ravi Strategist", "20 recognized foundations and one focused gap", "Global knowledge reuse prevents unnecessary repetition"],
        ],
        widths=[1.35, 3.15, 2.2],
    )
    heading(doc, "Demonstration narration", 2)
    panel_words(doc, "On this screen the recommendation is not hardcoded. The API rebuilt the learner context, checked every required prerequisite, sent only eligible candidates to the frozen Random Forest and stored the final reasons. I will now create new assessed evidence so you can see the path become stale and regenerate.")
    heading(doc, "If a service fails during the demo", 2)
    add_table(
        doc,
        ["Failure", "Expected behavior", "What to explain"],
        [
            ["PostgreSQL unavailable", "API readiness returns 503", "The app refuses to claim readiness without learner truth"],
            ["Inference artifact unavailable", "Inference returns 503", "There is no fabricated probability fallback"],
            ["Invalid feature request", "Inference returns 422", "Schema and bounds are part of the model contract"],
            ["NetworkX unavailable", "Prerequisite gates still work and graph bonus is omitted", "Safety stays in TypeScript; analytics degrade safely"],
            ["Path is stale", "Coordination excludes it until regeneration", "The app does not combine outdated recommendations"],
        ],
        widths=[1.8, 2.5, 2.4],
    )


QUESTIONS = [
    ("What is the main objective", "Determine the best safe next learning action from current learner evidence and make the reason visible.", "The pipeline combines global learner state, course targets, retention, confidence-aware prerequisite gates, candidate generation, checked Random Forest inference and versioned path construction."),
    ("What is your core novelty", "Cross-course learner knowledge and prerequisite graph intelligence drive an explainable adaptive path.", "The novelty is the integrated boundary: one global learner-skill state, evidence-bounded diagnostics, deterministic graph gates, NetworkX leverage, ML ranking only after eligibility and auditable regeneration."),
    ("Where exactly is machine learning used", "It ranks the safe candidate skills.", "FastAPI loads benefit-ranking-v2 and predicts benefit probability for each eligible 55-feature record. TypeScript owns mastery, retention, gates and path policy."),
    ("Why do you need ML if you have a graph", "The graph answers what is possible; ML estimates which possible option is most beneficial.", "Required edges create the feasible set. Random Forest orders that set from learner and candidate features. The graph contributes only a bounded structural bonus."),
    ("Can ML recommend a locked skill", "No.", "Locked candidates are removed before feature construction. constructPath also requires inference output to cover every and only eligible candidate."),
    ("Why Random Forest", "It won the declared validation ranking metric and handles nonlinear feature interactions.", "Selection used validation NDCG at 5. Random Forest reached 0.581021, slightly above the Highest Skill Gap baseline at 0.576408. The test result is mixed and is disclosed."),
    ("Is the model accurate", "The synthetic evaluation is reproducible, but real learner effectiveness is not yet proven.", "Test ROC-AUC is 0.801410 and Precision at 5 is 0.525333. Test NDCG at 5 is 0.578463, slightly below Highest Skill Gap at 0.581229."),
    ("How do you prevent data leakage", "The model uses only information available before the recommendation outcome.", "Identifiers and labels are separate, current outcome columns are blocked, outcome histories are shifted one interaction and learners do not overlap across train, validation and test."),
    ("Why synthetic data", "It allowed a reproducible ML pipeline without presenting invented records as real learners.", "The simulator is seed 42, labels come from an explicit benefit formula and the data remains outside PostgreSQL learner tables. Real validation remains future work."),
    ("How can fifteen questions evaluate a whole learner", "They cannot evaluate every skill completely. They gather the most useful starting evidence.", "The selector spends five slots on coverage, eight on confirmation and two on verification. Untested skills are returned as NOT TESTED and later evidence continues the estimate."),
    ("Can one question determine mastery", "No. One question produces PROBED.", "MASTERED needs at least three observations, a correct application or analysis item, target mastery and confidence at least 0.35."),
    ("What about a learner strong in advanced work but weak in basics", "LearnPath stores evidence per skill, so it can recognize advanced knowledge and repair the weak foundation.", "Diagnostic logic can identify FRAGILE FOUNDATION. The prerequisite graph exposes the weak edge and the candidate engine adds a supporting prerequisite without discarding advanced evidence."),
    ("What is confidence", "Confidence measures how much evidence supports mastery.", "It grows from count, source diversity, sessions, difficulty, consistency and retention evidence. Diagnostic-only evidence is capped at 0.45."),
    ("What is the difference between mastery and evidence state", "Mastery is the estimate; evidence state is the strength category behind it.", "UNKNOWN, ESTIMATED, ASSESSED and VERIFIED use explicit observation, session, consistency, confidence and diversity thresholds."),
    ("Why does lesson completion not update mastery", "Opening content is activity, not proof that the learner understood it.", "Resource events update activity and time only. Practice and assessment create server-scored evidence through the centralized mastery service."),
    ("How is forgetting implemented", "Retention decays from the last scored evidence while mastery remains the durable record.", "The exponential model adjusts its decay rate with confidence and evidence repetition. Revision needs earlier mastery at least 0.65 and current AT RISK or CRITICAL retention."),
    ("Why separate retention from mastery", "It distinguishes never learned from learned but currently forgotten.", "A low mastery gap becomes LEARN. Strong earlier mastery with decayed retention becomes REVISION. These are different candidate kinds and explanations."),
    ("How does the prerequisite graph work", "Each required edge says which foundation and threshold a skill needs.", "The service recursively loads the context graph, validates dependency order and compares confidence-adjusted global mastery with the threshold on every REQUIRED edge."),
    ("Why NetworkX", "It calculates graph structure that is difficult to express as ordinary table joins.", "It produces topological layers, descendants, betweenness, shortest foundation routes, bottlenecks and counterfactual unlocks from a read-only projection."),
    ("Why not Neo4j", "The graph is small and PostgreSQL already owns the curriculum transactionally.", "A second graph database would duplicate truth and require synchronization. Neo4j would be justified by much larger graph-native workloads or heterogeneous relationship discovery."),
    ("How is the path personalized", "Its order depends on this learner's evidence, retention, course targets and safe model scores.", "The stored items include mastery, confidence, evidence state, retention, prerequisites, benefit probability, bounded bonuses, reason codes and model or policy versions."),
    ("Why is path regeneration explicit", "The learner should know when the recommendation changed.", "New evidence marks an ACTIVE path STALE. A transaction supersedes it and creates a new ACTIVE version plus change summary."),
    ("How do multiple courses work", "Every course has its own progress and path, but all courses read the same Skill Passport.", "Coordination groups the same Learn Next skill across current active paths and adds bounded shared-course and goal relevance bonuses without merging paths."),
    ("How do you explain a recommendation", "Show the gap, prerequisites, retention, model benefit and course context.", "The path persists explanation text and reason codes such as ML HIGHEST PRIORITY, PREREQUISITES SATISFIED, MASTERY GAP, RETENTION REVISION and NETWORKX GATEWAY PRIORITY."),
    ("How do you trust the inference service", "It refuses incompatible data or artifacts.", "Startup verifies SHA-256, versions, feature count, exact order and positive class. Invalid requests return 422; missing or incompatible artifacts return 503."),
    ("What happens if NetworkX is down", "The safe prerequisite result still works.", "The TypeScript engine remains authoritative, the graph analytics status becomes safe mode and the gateway bonus is omitted."),
    ("How do you prevent fake improvement claims", "We call later outcomes observed, not caused.", "Attribution uses a 30-day window, a stored baseline and a 0.05 threshold. Missing baseline becomes OBSERVED NO BASELINE."),
    ("How do you handle small evaluation samples", "We withhold unstable conclusions.", "Governance requires 30 responses for rates, 30 assessed outcomes for calibration, 50 plus 50 predictions for drift and 100 assessed outcomes for retraining eligibility."),
    ("What would you improve next", "Calibrate with real learner data and evaluate learning gain prospectively.", "Next steps are item calibration, expert graph review, controlled learner study, subgroup and calibration analysis, and human-reviewed model promotion."),
    ("What does the frontend calculate", "It does not calculate learner intelligence.", "React renders typed API responses and sends learner actions. Mastery, retention, gates, scores and paths are created by backend services."),
]


def question_bank(doc: Document):
    heading(doc, "11 Reviewer question bank", 1, new_page=True)
    para(doc, "Use the short answer first. Continue with the technical answer only when the reviewer asks how the claim is implemented.")
    for index, (question, simple, technical) in enumerate(QUESTIONS, start=1):
        if index in (7, 13, 19, 25):
            doc.add_page_break()
        heading(doc, f"Question {index} {question}", 2)
        para(doc, f"Short answer  {simple}", lead="Short answer  ")
        para(doc, f"Technical answer  {technical}", lead="Technical answer  ")


def glossary(doc: Document):
    heading(doc, "12 Glossary and final revision sheet", 1, new_page=True)
    add_table(
        doc,
        ["Term", "Meaning in LearnPath"],
        [
            ["Global skill state", "One learner-specific knowledge record reused across all courses"],
            ["Mastery", "Zero-to-one estimate of durable demonstrated knowledge"],
            ["Confidence", "Zero-to-one strength of evidence supporting mastery"],
            ["Evidence state", "UNKNOWN, ESTIMATED, ASSESSED or VERIFIED support category"],
            ["Retention", "Current projection of how accessible earlier knowledge remains"],
            ["Evidence", "An immutable scored performance observation"],
            ["Course target mastery", "The mastery level this course expects for a skill"],
            ["Required prerequisite", "A blocking edge with a mastery threshold"],
            ["Recommended prerequisite", "Advisory preparation that does not block"],
            ["Candidate", "A relevant skill considered for learning, revision or support"],
            ["Eligible", "Candidate with every required prerequisite satisfied"],
            ["Locked", "Candidate with at least one failed required edge"],
            ["Excluded", "Candidate already satisfied and not due for revision"],
            ["Benefit probability", "Random Forest estimate for an eligible candidate"],
            ["Gateway score", "Bounded NetworkX structural leverage measure"],
            ["Learn Next", "The highest-priority actionable item in the persisted path"],
            ["Stale path", "Current path whose decision inputs changed after new evidence"],
            ["Superseded path", "Immutable former path kept for comparison"],
            ["Synthetic data", "Simulated data clearly separated from real learner records"],
            ["Governance gate", "Minimum sample requirement before a metric or action is allowed"],
        ],
        widths=[2.0, 4.7],
        size=8.7,
    )
    heading(doc, "Final revision sheet", 2)
    add_table(
        doc,
        ["Reviewer asks", "Remember this answer"],
        [
            ["What is LearnPath", "An evidence-driven system for choosing and explaining the learner's next safe skill"],
            ["Where is ML", "Only in benefit ranking of eligible candidates"],
            ["Where is the knowledge graph", "PostgreSQL graph truth, TypeScript gates and NetworkX analytics"],
            ["Why can one answer not prove mastery", "One answer is PROBED; mastery needs repeated varied application evidence"],
            ["How are multiple courses handled", "Global knowledge is shared; course progress and paths remain separate"],
            ["How does the path adapt", "New evidence makes it stale; regeneration creates a new auditable version"],
            ["Why trust it", "Immutable evidence, conservative gates, checked artifacts, reproducible evaluation and honest limitations"],
            ["What remains to prove", "Improved learning outcomes with real learners"],
        ],
        widths=[2.2, 4.5],
    )
    panel_words(doc, "LearnPath makes the next-learning decision explainable from evidence. Its engineering controls are implemented and testable. Its educational effectiveness is the next empirical research question.")


def build():
    create_equation_images()
    create_handbook_diagrams()
    doc = Document()
    configure(doc)
    cover(doc)
    contents(doc)
    presentation_scripts(doc)
    plain_system(doc)
    architecture(doc)
    knowledge_evidence(doc)
    diagnostic_graph(doc)
    recommendation_adaptation(doc)
    supporting_modules(doc)
    numerical_examples(doc)
    trust_validation(doc)
    demo_guide(doc)
    question_bank(doc)
    glossary(doc)
    properties = doc.core_properties
    properties.title = "LearnPath Reviewer Explanation Handbook"
    properties.subject = "Technical and nontechnical explanations for presenting and defending LearnPath"
    properties.author = "LearnPath Project Team"
    properties.keywords = "LearnPath reviewer handbook adaptive learning knowledge graph diagnostics machine learning"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
